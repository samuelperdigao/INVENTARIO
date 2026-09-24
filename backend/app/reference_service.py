"""Importação segura e conciliação de referências de lotes.

O serviço aceita somente ``.xlsx``. O workbook fica apenas em memória durante
o processamento e nenhuma coluna, célula ou arquivo original é persistido.
Depois da confirmação, a única carga salva é o conjunto de números de lote.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
import math
from pathlib import PurePath
import re
import zipfile
from typing import Any, Iterable
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.lot_rules import LOT_VALIDATION_MESSAGE, is_valid_lot, validate_lot
from app.persistence import InventoryEntryRow, InventoryReferenceRow, ReferenceLotRow, UserRow


SOURCE_TYPE = "SAP_EXCEL"
ACTIVE_STATUS = "ACTIVE"
REMOVED_STATUS = "REMOVED"
MAX_COLUMNS = 256
HEADER_SCAN_ROWS = 25
WARNING_LIMIT = 20
LOT_HEADERS = frozenset({"lote", "lotes"})
MISSING_LOT_HEADER_MESSAGE = "Não foi possível localizar a coluna ‘Lote’ ou ‘Lotes’ na planilha do SAP. Confira o arquivo selecionado."


class ReferenceImportError(ValueError):
    """Erro de entrada que deve ser devolvido como validação HTTP 422."""


@dataclass(frozen=True)
class ReferenceColumn:
    index: int
    label: str


@dataclass(frozen=True)
class ParsedReference:
    original_filename: str
    header_row: int | None
    columns: tuple[ReferenceColumn, ...]
    selected_column: int | None
    selected_column_label: str | None
    total_rows: int
    valid_lot_occurrences: int
    unique_lots: int
    duplicate_rows: int
    ignored_rows: int
    lots: tuple[str, ...]
    sample: tuple[str, ...]
    warnings: tuple[str, ...]

    @property
    def requires_column_selection(self) -> bool:
        return self.selected_column is None

    def preview_payload(self) -> dict[str, Any]:
        return {
            "originalFilename": self.original_filename,
            "sourceType": SOURCE_TYPE,
            "headerRow": self.header_row,
            "columns": [{"index": column.index, "label": column.label} for column in self.columns],
            "selectedColumn": self.selected_column,
            "selectedColumnLabel": self.selected_column_label,
            "requiresColumnSelection": self.requires_column_selection,
            "totalRows": self.total_rows,
            "validLotOccurrences": self.valid_lot_occurrences,
            "uniqueLots": self.unique_lots,
            "duplicateRows": self.duplicate_rows,
            "ignoredRows": self.ignored_rows,
            "sample": list(self.sample),
            "warnings": list(self.warnings),
        }


def safe_original_filename(filename: str | None) -> str:
    """Remove caminhos e caracteres de controle do nome exibido no card."""

    raw = (filename or "referencia.xlsx").replace("\x00", "")
    basename = re.split(r"[\\/]", raw)[-1].strip()
    basename = re.sub(r"[\x00-\x1f\x7f]+", "_", basename)
    return (basename or "referencia.xlsx")[:255]


def _normalized_header(value: object) -> str:
    return str(value or "").strip().casefold()


def _column_label(value: object, index: int) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip())
    return text[:120] or f"Coluna {get_column_letter(index)}"


def _safe_zip_bytes(content: bytes, max_bytes: int) -> None:
    if len(content) > max_bytes:
        raise ReferenceImportError(f"O arquivo excede o limite de {max_bytes // (1024 * 1024)} MB.")
    if not content or not zipfile.is_zipfile(BytesIO(content)):
        raise ReferenceImportError("O arquivo não é uma planilha Excel .xlsx válida.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ReferenceImportError("O arquivo não possui a estrutura interna esperada de um .xlsx.")
            if any(
                name.startswith(("/", "\\")) or ".." in PurePath(name).parts
                for name in names
            ):
                raise ReferenceImportError("A planilha contém um caminho interno inválido.")
            if any("vbaProject" in name or name.startswith("xl/externalLinks/") for name in names):
                raise ReferenceImportError("Macros e links externos não são aceitos na referência.")
            if sum(info.file_size for info in archive.infolist()) > max_bytes * 20:
                raise ReferenceImportError("A planilha compactada excede o limite seguro de expansão.")
            if archive.testzip() is not None:
                raise ReferenceImportError("A planilha está corrompida.")
    except zipfile.BadZipFile as error:
        raise ReferenceImportError("O arquivo não é uma planilha Excel .xlsx válida.") from error


def _is_nonempty(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _load_single_data_sheet(content: bytes) -> tuple[Any, Any, list[list[Any]]]:
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as error:
        raise ReferenceImportError("Não foi possível abrir a planilha Excel com segurança.") from error

    try:
        data_sheets: list[tuple[Any, list[list[Any]]]] = []
        for worksheet in workbook.worksheets:
            preview_rows: list[list[Any]] = []
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=HEADER_SCAN_ROWS,
                min_col=1,
                max_col=MAX_COLUMNS,
            ):
                values = [cell.value for cell in row]
                preview_rows.append(values)
            if any(_is_nonempty(value) for row in preview_rows for value in row):
                data_sheets.append((worksheet, preview_rows))
        if not data_sheets:
            raise ReferenceImportError("A planilha Excel está vazia.")
        if len(data_sheets) > 1:
            raise ReferenceImportError("A planilha possui mais de uma aba com dados. Envie uma única aba exportada do SAP.")
        worksheet, preview_rows = data_sheets[0]
        return workbook, worksheet, preview_rows
    except ReferenceImportError:
        workbook.close()
        raise


def _detect_columns(preview_rows: list[list[Any]]) -> tuple[tuple[ReferenceColumn, ...], int | None, int | None]:
    nonempty_columns = {
        column_index
        for row in preview_rows
        for column_index, value in enumerate(row, start=1)
        if _is_nonempty(value)
    }
    if not nonempty_columns:
        raise ReferenceImportError("A planilha Excel está vazia.")

    header_candidates: list[tuple[int, int]] = []
    for row_index, row in enumerate(preview_rows, start=1):
        for column_index, value in enumerate(row, start=1):
            if _normalized_header(value) in LOT_HEADERS:
                header_candidates.append((row_index, column_index))

    if not header_candidates:
        raise ReferenceImportError(MISSING_LOT_HEADER_MESSAGE)
    candidate_columns = {candidate[1] for candidate in header_candidates}
    if len(candidate_columns) != 1:
        raise ReferenceImportError("A planilha deve conter uma única coluna com cabeçalho ‘Lote’ ou ‘Lotes’.")
    header_row, automatic_column = min(header_candidates)

    columns = tuple(
        ReferenceColumn(
            index=index,
            label=_column_label(
                preview_rows[header_row - 1][index - 1] if header_row and index <= len(preview_rows[header_row - 1]) else None,
                index,
            ),
        )
        for index in sorted(nonempty_columns)
    )
    return columns, automatic_column, header_row


def _zero_format_width(number_format: str | None) -> int | None:
    pattern = str(number_format or "").split(";", 1)[0]
    if re.search(r"[eE][+-]?", pattern):
        return None
    pattern = re.sub(r'"[^"\r\n]*"|\[[^\]]*\]|_[^\r\n]|\*[^\r\n]', "", pattern)
    integer_part = pattern.split(".", 1)[0]
    placeholders = re.findall(r"[0#?]", integer_part)
    if not placeholders or "0" not in placeholders:
        return None
    return len(placeholders)


def _decimal_to_lot(value: Decimal, number_format: str | None, *, numeric_warning: bool) -> tuple[str | None, str | None]:
    if not value.is_finite() or value != value.to_integral_value():
        return None, "A célula numérica não representa um número inteiro de lote."
    text = format(value, "f").split(".", 1)[0]
    width = _zero_format_width(number_format)
    if width and len(text) < width:
        text = text.zfill(width)
    warning = "A célula numérica não permite recuperar com segurança eventuais zeros à esquerda."
    if width:
        warning = None
    if numeric_warning and abs(value) > Decimal(2**53):
        warning = "A célula numérica pode ter perdido precisão antes da importação pelo Excel."
    return text, warning


def _normalize_cell(cell: Any) -> tuple[str | None, str | None]:
    value = cell.value
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if getattr(cell, "data_type", None) == "f":
        return None, "A linha com fórmula foi ignorada; informe o lote como valor ou texto."
    if isinstance(value, bool):
        return None, "A célula booleana foi ignorada."
    if isinstance(value, str):
        text = value.strip().replace("\u200b", "").replace("\u200c", "").replace("\u200d", "").replace("\ufeff", "")
        if not text:
            return None, None
        if re.fullmatch(r"\d+", text):
            return text, None
        if re.fullmatch(r"[+-]?\d+(?:[\.,]\d+)?[eE][+-]?\d+", text):
            try:
                decimal_value = Decimal(text.replace(",", "."))
            except InvalidOperation:
                decimal_value = None
            if decimal_value is not None:
                return _decimal_to_lot(decimal_value, getattr(cell, "number_format", None), numeric_warning=False)
        return None, "A célula contém caracteres que não formam um número de lote."
    if isinstance(value, int):
        return _decimal_to_lot(Decimal(value), getattr(cell, "number_format", None), numeric_warning=False)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None, "A célula numérica não é finita."
        return _decimal_to_lot(Decimal(str(value)), getattr(cell, "number_format", None), numeric_warning=True)
    if isinstance(value, Decimal):
        return _decimal_to_lot(value, getattr(cell, "number_format", None), numeric_warning=True)
    return None, "A célula contém um tipo não suportado para número de lote."


def parse_xlsx_reference(
    content: bytes,
    filename: str | None,
    *,
    selected_column: int | None = None,
    max_bytes: int,
    max_rows: int,
) -> ParsedReference:
    """Valida e transforma o workbook em uma lista deduplicada de lotes."""

    safe_filename = safe_original_filename(filename)
    if not safe_filename.casefold().endswith(".xlsx"):
        raise ReferenceImportError("Envie um arquivo Excel .xlsx exportado do SAP.")
    _safe_zip_bytes(content, max_bytes)
    workbook, worksheet, preview_rows = _load_single_data_sheet(content)
    try:
        columns, automatic_column, header_row = _detect_columns(preview_rows)
        selected = selected_column if selected_column is not None else automatic_column
        available_columns = {column.index: column.label for column in columns}
        if selected is not None and (selected < 1 or selected > MAX_COLUMNS or selected not in available_columns):
            raise ReferenceImportError("Selecione uma coluna disponível na prévia da planilha.")
        if selected != automatic_column:
            raise ReferenceImportError("A importação da referência SAP utiliza exclusivamente a coluna com cabeçalho ‘Lote’ ou ‘Lotes’.")
        selected_label = available_columns.get(selected) if selected is not None else None
        first_data_row = header_row + 1 if header_row else 1
        total_rows = 0
        valid_occurrences = 0
        ignored_rows = 0
        duplicate_rows = 0
        unique_lots: list[str] = []
        seen: set[str] = set()
        warnings: list[str] = []
        for row in worksheet.iter_rows(min_row=first_data_row, min_col=selected, max_col=selected):
            total_rows += 1
            if total_rows > max_rows:
                raise ReferenceImportError(f"A planilha excede o limite de {max_rows:,} linhas processáveis.")
            lot, warning = _normalize_cell(row[0])
            if warning and len(warnings) < WARNING_LIMIT:
                warnings.append(f"Linha {first_data_row + total_rows - 1}: {warning}")
            if lot is None:
                ignored_rows += 1
                continue
            if not is_valid_lot(lot):
                ignored_rows += 1
                if len(warnings) < WARNING_LIMIT:
                    warnings.append(f"Linha {first_data_row + total_rows - 1}: {LOT_VALIDATION_MESSAGE}")
                continue
            lot = validate_lot(lot)
            valid_occurrences += 1
            if lot in seen:
                duplicate_rows += 1
                continue
            seen.add(lot)
            unique_lots.append(lot)
        if not unique_lots:
            warnings.append(f"Nenhum lote válido foi encontrado na coluna ‘{selected_label}’. A referência não pode ser confirmada.")
        if duplicate_rows and len(warnings) < WARNING_LIMIT:
            warnings.append(f"{duplicate_rows} ocorrência(s) repetida(s) foram mantidas apenas uma vez.")
        return ParsedReference(
            original_filename=safe_filename,
            header_row=header_row,
            columns=columns,
            selected_column=selected,
            selected_column_label=selected_label,
            total_rows=total_rows,
            valid_lot_occurrences=valid_occurrences,
            unique_lots=len(unique_lots),
            duplicate_rows=duplicate_rows,
            ignored_rows=ignored_rows,
            lots=tuple(unique_lots),
            sample=tuple(unique_lots[:5]),
            warnings=tuple(warnings),
        )
    finally:
        workbook.close()


def active_reference(session: Session, inventory_id: str, *, for_update: bool = False) -> InventoryReferenceRow | None:
    statement = select(InventoryReferenceRow).where(
        InventoryReferenceRow.inventory_id == inventory_id,
        InventoryReferenceRow.status == ACTIVE_STATUS,
    )
    if for_update:
        statement = statement.with_for_update()
    return session.scalar(statement)


def _iso(value: datetime | None) -> str | None:
    return value.isoformat() if value else None


def reference_metadata(session: Session, reference: InventoryReferenceRow) -> dict[str, Any]:
    creator = session.get(UserRow, reference.created_by_user_id) if reference.created_by_user_id else None
    return {
        "sourceType": reference.source_type,
        "originalFilename": reference.original_filename,
        "importedAt": _iso(reference.imported_at),
        "updatedAt": _iso(reference.updated_at),
        "totalLots": reference.total_lots,
        "revision": reference.revision,
        "createdByName": creator.display_name if creator else None,
    }


def report_reference_metadata(session: Session, inventory_id: str) -> dict[str, Any] | None:
    reference = active_reference(session, inventory_id)
    if reference is None:
        return None
    return reference_metadata(session, reference)


def reference_lot_numbers(session: Session, reference_id: str) -> tuple[str, ...]:
    return tuple(session.scalars(select(ReferenceLotRow.lot_number).where(ReferenceLotRow.reference_id == reference_id)))


def replace_reference(
    session: Session,
    *,
    inventory_id: str,
    filename: str,
    lots: Iterable[str],
    user_id: str,
) -> InventoryReferenceRow:
    try:
        normalized_lots = tuple(dict.fromkeys(validate_lot(lot) for lot in lots))
    except ValueError as error:
        raise ReferenceImportError(str(error)) from error
    if not normalized_lots:
        raise ReferenceImportError("Nenhum número de lote válido foi encontrado.")
    now = datetime.now(timezone.utc)
    reference = active_reference(session, inventory_id, for_update=True)
    if reference is None:
        reference = session.scalar(
            select(InventoryReferenceRow).where(InventoryReferenceRow.inventory_id == inventory_id).with_for_update()
        )
    if reference is None:
        reference = InventoryReferenceRow(
            id=str(uuid4()),
            inventory_id=inventory_id,
            source_type=SOURCE_TYPE,
            status=ACTIVE_STATUS,
            original_filename=filename,
            imported_at=now,
            updated_at=now,
            removed_at=None,
            created_by_user_id=user_id,
            total_lots=len(normalized_lots),
            revision=1,
        )
        session.add(reference)
        session.flush()
    else:
        reference.source_type = SOURCE_TYPE
        reference.status = ACTIVE_STATUS
        reference.original_filename = filename
        reference.imported_at = now
        reference.updated_at = now
        reference.removed_at = None
        reference.created_by_user_id = user_id
        reference.total_lots = len(normalized_lots)
        reference.revision += 1
        session.execute(delete(ReferenceLotRow).where(ReferenceLotRow.reference_id == reference.id))
    session.add_all(
        ReferenceLotRow(id=str(uuid4()), reference_id=reference.id, lot_number=lot)
        for lot in normalized_lots
    )
    session.flush()
    return reference


def remove_reference(session: Session, inventory_id: str) -> InventoryReferenceRow | None:
    reference = active_reference(session, inventory_id, for_update=True)
    if reference is None:
        return None
    now = datetime.now(timezone.utc)
    reference.status = REMOVED_STATUS
    reference.updated_at = now
    reference.removed_at = now
    reference.total_lots = 0
    reference.revision += 1
    session.execute(delete(ReferenceLotRow).where(ReferenceLotRow.reference_id == reference.id))
    session.flush()
    return reference


def _physical_stats(session: Session, inventory_id: str) -> dict[str, dict[str, Any]]:
    stats: dict[str, dict[str, Any]] = {}
    rows = session.execute(
        select(
            InventoryEntryRow.lot,
            InventoryEntryRow.side,
            InventoryEntryRow.bay,
            InventoryEntryRow.layer,
            InventoryEntryRow.quantity,
        ).where(
            InventoryEntryRow.inventory_id == inventory_id,
            InventoryEntryRow.tombstone.is_(False),
        )
    )
    for lot, side, bay, layer, quantity in rows:
        item = stats.setdefault(lot, {"quantity": 0, "occurrences": 0, "locations": set()})
        item["quantity"] += int(quantity)
        item["occurrences"] += 1
        item["locations"].add((side, bay, layer or ""))
    return stats


def reference_state(
    session: Session,
    inventory_id: str,
    *,
    page: int = 1,
    page_size: int = 50,
    query: str = "",
) -> dict[str, Any]:
    reference = active_reference(session, inventory_id)
    if reference is None:
        return {
            "reference": None,
            "summary": {
                "available": False,
                "totalLots": 0,
                "foundLots": 0,
                "pendingLots": 0,
                "outsideReferenceLots": 0,
                "fragmentedLots": 0,
                "physicalDistinctLots": 0,
            },
            "lots": [],
            "page": 1,
            "pageSize": page_size,
            "totalMatchingLots": 0,
            "totalPages": 0,
        }

    physical = _physical_stats(session, inventory_id)
    reference_lots = set(reference_lot_numbers(session, reference.id))
    physical_lots = set(physical)
    found_lots = reference_lots & physical_lots
    clean_query = query.strip()
    statement = select(ReferenceLotRow).where(ReferenceLotRow.reference_id == reference.id)
    count_statement = select(func.count()).select_from(ReferenceLotRow).where(ReferenceLotRow.reference_id == reference.id)
    if clean_query:
        if not clean_query.isdigit():
            return {
                "reference": reference_metadata(session, reference),
                "summary": {
                    "available": True,
                    "totalLots": len(reference_lots),
                    "foundLots": len(found_lots),
                    "pendingLots": len(reference_lots - physical_lots),
                    "outsideReferenceLots": len(physical_lots - reference_lots),
                    "fragmentedLots": sum(len(item["locations"]) > 1 for item in physical.values()),
                    "physicalDistinctLots": len(physical_lots),
                },
                "lots": [],
                "page": 1,
                "pageSize": page_size,
                "totalMatchingLots": 0,
                "totalPages": 0,
            }
        pattern = f"%{clean_query}%"
        statement = statement.where(ReferenceLotRow.lot_number.like(pattern))
        count_statement = count_statement.where(ReferenceLotRow.lot_number.like(pattern))
    total_matching = int(session.scalar(count_statement) or 0)
    total_pages = (total_matching + page_size - 1) // page_size if total_matching else 0
    current_page = min(max(page, 1), total_pages or 1)
    rows = session.scalars(
        statement.order_by(ReferenceLotRow.lot_number).offset((current_page - 1) * page_size).limit(page_size)
    ).all()
    lot_items = [
        {
            "lotNumber": row.lot_number,
            "foundPhysically": row.lot_number in physical,
            "physicalQuantity": int(physical.get(row.lot_number, {}).get("quantity", 0)),
            "physicalOccurrences": int(physical.get(row.lot_number, {}).get("occurrences", 0)),
            "fragmented": len(physical.get(row.lot_number, {}).get("locations", set())) > 1,
        }
        for row in rows
    ]
    return {
        "reference": reference_metadata(session, reference),
        "summary": {
            "available": True,
            "totalLots": len(reference_lots),
            "foundLots": len(found_lots),
            "pendingLots": len(reference_lots - physical_lots),
            "outsideReferenceLots": len(physical_lots - reference_lots),
            "fragmentedLots": sum(len(item["locations"]) > 1 for item in physical.values()),
            "physicalDistinctLots": len(physical_lots),
        },
        "lots": lot_items,
        "page": current_page,
        "pageSize": page_size,
        "totalMatchingLots": total_matching,
        "totalPages": total_pages,
    }


def reference_match(session: Session, inventory_id: str, lot: str) -> dict[str, Any]:
    lot = validate_lot(lot)
    reference = active_reference(session, inventory_id)
    if reference is None:
        return {"referenceAvailable": False, "lot": lot, "inReference": None}
    found = session.scalar(
        select(ReferenceLotRow.id).where(
            ReferenceLotRow.reference_id == reference.id,
            ReferenceLotRow.lot_number == lot,
        )
    ) is not None
    return {"referenceAvailable": True, "lot": lot, "inReference": found}
