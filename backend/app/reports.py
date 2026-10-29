"""Modelo único e exportadores oficiais dos relatórios V1.

As regras de classificação continuam no motor de análise. Este módulo apenas
prepara uma representação comum e a renderiza em formatos diferentes.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone, tzinfo
from io import BytesIO
from typing import Any, Iterable, Mapping
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import xlwt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import CondPageBreak, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.engine import AnalysisEntry, analyze_entries
from app.lot_rules import validate_lot
from app.presentation import apply_report_presentation, build_lot_presentation


BLUE_DARK = "1F4E78"
BLUE_LIGHT = "D9EAF7"
BLUE_PALE = "EDF4FA"
GREY_LIGHT = "F2F2F2"
GREY_BORDER = "B8C4CE"
WHITE = "FFFFFF"
ALERT_RED = "F4CCCC"
ALERT_ORANGE = "FCE4D6"
ALERT_YELLOW = "FFF2CC"
GOOD_GREEN = "E2F0D9"

_ALERT_COLORS = {
    "single-piece": ALERT_YELLOW,
    "multiple-pieces": ALERT_ORANGE,
    "distributed": ALERT_RED,
    "review": ALERT_ORANGE,
    "ok": GOOD_GREEN,
}


def _natural_key(value: str) -> list[object]:
    import re

    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def build_inventory_report_data(
    inventory_id: str,
    inventory_date: str,
    revision: int,
    entries: Iterable[AnalysisEntry],
    generated_at: datetime | None = None,
    *,
    reference_lots: Iterable[str] | None = None,
    reference_metadata: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Prepara a fonte única consumida por todos os exportadores."""

    raw_entries = list(entries)
    analysis = analyze_entries(inventory_id, revision, raw_entries)
    if reference_lots is not None:
        normalized_reference = {validate_lot(lot) for lot in reference_lots}
        physical_lots = {str(lot["lot"]) for lot in analysis["lots"]}
        for lot in analysis["lots"]:
            lot["referenceStatus"] = "EXPECTED_FOUND" if lot["lot"] in normalized_reference else "OUTSIDE_REFERENCE"
            lot["physicalFound"] = True
        for missing_lot in sorted(normalized_reference - physical_lots, key=_natural_key):
            analysis["lots"].append(
                {
                    "lot": missing_lot,
                    "totalQuantity": 0,
                    "locations": [],
                    "fragmented": False,
                    "classification": "REVISAR",
                    "primaryLocation": None,
                    "displacedQuantity": 0,
                    "recommendation": "Nenhum lançamento físico foi encontrado para este lote previsto.",
                    "referenceStatus": "EXPECTED_MISSING",
                    "physicalFound": False,
                }
            )
        analysis["lots"].sort(key=lambda lot: _natural_key(str(lot["lot"])))
        physical_fragmented = sum(1 for lot in analysis["lots"] if lot.get("physicalFound") and lot.get("fragmented"))
        analysis["reference"] = {
            "available": True,
            **dict(reference_metadata or {}),
            "totalLots": len(normalized_reference),
            "foundLots": len(normalized_reference & physical_lots),
            "missingLots": len(normalized_reference - physical_lots),
            "outsideReferenceLots": len(physical_lots - normalized_reference),
            "fragmentedLots": physical_fragmented,
            "physicalDistinctLots": len(physical_lots),
        }
    analysis = apply_report_presentation(analysis)
    analysis["generatedAt"] = (generated_at or datetime.now().astimezone()).isoformat()
    records = [
        {
            "side": entry.side,
            "bay": entry.bay.strip(),
            "layer": entry.layer.strip() if entry.layer else None,
            "lot": validate_lot(entry.lot),
            "quantity": entry.quantity,
        }
        for entry in raw_entries
    ]
    records.sort(
        key=lambda record: (
            0 if record["side"] == "EF" else 1,
            _natural_key(record["bay"]),
            _natural_key(record["layer"] or ""),
            _natural_key(record["lot"]),
        )
    )
    analysis["inventoryDate"] = inventory_date
    analysis["records"] = records
    analysis["totalPieces"] = sum(record["quantity"] for record in records)
    analysis["totalRecords"] = len(records)
    return analysis


def build_consolidated_report(
    inventory_id: str,
    inventory_date: str,
    revision: int,
    entries: Iterable[AnalysisEntry],
    generated_at: datetime | None = None,
    *,
    reference_lots: Iterable[str] | None = None,
    reference_metadata: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Compatibilidade para consumidores que usam o nome histórico."""

    return build_inventory_report_data(
        inventory_id,
        inventory_date,
        revision,
        entries,
        generated_at,
        reference_lots=reference_lots,
        reference_metadata=reference_metadata,
    )


REPORT_TIMEZONE_NAME = "America/Sao_Paulo"
NO_CONFERENCE_ROW: list[object] = ["—", None, "Nenhum lote para conferência", "—", "—", None, "Nenhum lote precisa de conferência."]
INVENTORY_HEADERS = ["Lado", "Vão", "Camada", "Lote", "Quantidade de peças"]
CONSOLIDATED_HEADERS = ["Lote", "Total de peças", "Localização", "Situação"]
CONFERENCE_HEADERS = ["Lote", "Total", "Situação", "Local principal", "Outros locais", "Peças fora", "Ação recomendada"]
CONCILIATION_HEADERS = ["Lote", "Referência", "Físico", "Localização", "Qtd. física", "Condição"]


class ReportConsistencyError(ValueError):
    """Os totais do relatório não fecham; nenhum arquivo deve ser gerado."""


def _report_timezone() -> tzinfo:
    try:
        return ZoneInfo(REPORT_TIMEZONE_NAME)
    except ZoneInfoNotFoundError:
        # Sem base IANA no host, vale o fuso oficial de Brasília sem horário de verão.
        return timezone(timedelta(hours=-3), "BRT")


def format_generated_at(value: object) -> str:
    """Converte o instante de geração para o horário de Brasília legível."""

    if isinstance(value, datetime):
        moment = value
    elif isinstance(value, str) and value:
        try:
            moment = datetime.fromisoformat(value)
        except ValueError:
            return value
    else:
        return "Não informado"
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    local = moment.astimezone(_report_timezone())
    return f"{local:%d/%m/%Y} às {local:%H:%M}"


def parse_report_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def format_report_date(value: object) -> str:
    parsed = parse_report_date(value)
    return f"{parsed:%d/%m/%Y}" if parsed else str(value)


def format_int(value: int) -> str:
    """Formata inteiros com separador de milhar brasileiro (2.667)."""

    return f"{int(value):,}".replace(",", ".")


def _bay_value(bay: str) -> object:
    """Vão vira número no Excel somente quando a conversão é exata."""

    return int(bay) if bay.isdigit() and str(int(bay)) == bay else bay


def _presentation(lot: Mapping[str, Any]) -> Mapping[str, Any]:
    return lot.get("presentation") or build_lot_presentation(lot)


def _has_reference(report: Mapping[str, Any]) -> bool:
    reference = report.get("reference")
    return isinstance(reference, Mapping) and bool(reference.get("available"))


def _is_physical(lot: Mapping[str, Any]) -> bool:
    return bool(lot.get("physicalFound", lot.get("referenceStatus") != "EXPECTED_MISSING"))


def verify_report_totals(report: Mapping[str, Any]) -> None:
    """Confere que lançamentos, lotes, resumo e subtotais por vão fecham.

    Qualquer divergência interrompe a exportação: é preferível não entregar
    um arquivo a entregar um arquivo com totais que não batem.
    """

    records = list(report.get("records") or [])
    lots = [lot for lot in report.get("lots", []) if isinstance(lot, Mapping)]
    physical_lots = [lot for lot in lots if _is_physical(lot)]
    summary = report.get("summary") or {}
    problems: list[str] = []

    records_pieces = sum(int(record["quantity"]) for record in records)
    lots_pieces = sum(int(lot.get("totalQuantity", 0)) for lot in physical_lots)
    summary_pieces = int(report.get("totalPieces", -1))
    bay_pieces = sum(group["pieces"] for group in _bay_groups(records))
    if not records_pieces == lots_pieces == summary_pieces == bay_pieces:
        problems.append(
            "peças não fecham: lançamentos "
            f"{format_int(records_pieces)}, lotes {format_int(lots_pieces)}, resumo {format_int(summary_pieces)}, "
            f"subtotais por vão {format_int(bay_pieces)}"
        )
    if len(records) != int(report.get("totalRecords", -1)):
        problems.append(f"registros não fecham: {len(records)} lançamentos para {report.get('totalRecords')} no resumo")

    pieces_by_lot: dict[str, int] = {}
    for record in records:
        pieces_by_lot[str(record["lot"])] = pieces_by_lot.get(str(record["lot"]), 0) + int(record["quantity"])
    lots_by_id = {str(lot["lot"]): int(lot.get("totalQuantity", 0)) for lot in physical_lots}
    if pieces_by_lot != lots_by_id:
        divergent = sorted(set(pieces_by_lot) ^ set(lots_by_id) | {
            lot for lot in set(pieces_by_lot) & set(lots_by_id) if pieces_by_lot[lot] != lots_by_id[lot]
        })
        problems.append(f"lotes divergentes dos lançamentos: {', '.join(divergent[:5])}")
    for lot in physical_lots:
        located = sum(int(location.get("quantity", 0)) for location in lot.get("locations", []) if isinstance(location, Mapping))
        if located != int(lot.get("totalQuantity", 0)):
            problems.append(f"lote {lot['lot']} soma {located} nos locais e {lot.get('totalQuantity')} no total")
    if int(summary.get("lotsAnalyzed", -1)) != len(physical_lots):
        problems.append(f"total de lotes do resumo ({summary.get('lotsAnalyzed')}) difere dos {len(physical_lots)} lotes analisados")
    if "lotsOk" in summary and int(summary["lotsOk"]) + int(summary.get("lotsForConference", 0)) != len(lots):
        problems.append("lotes OK somados aos lotes para conferência não fecham com o total de lotes")
    if problems:
        raise ReportConsistencyError(
            "Relatório não gerado: os totais não conferem (" + "; ".join(problems) + ")."
        )


def _bay_groups(records: list[Mapping[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str], dict[str, Any]] = {}
    for record in records:
        key = (str(record["side"]), str(record["bay"]))
        group = groups.setdefault(key, {"side": key[0], "bay": key[1], "records": 0, "pieces": 0})
        group["records"] += 1
        group["pieces"] += int(record["quantity"])
    return list(groups.values())


@dataclass(frozen=True)
class ReportView:
    """Apresentação comum a todos os exportadores.

    Só formata e organiza o que `build_inventory_report_data` já calculou;
    nenhum exportador recalcula totais ou classificações.
    """

    inventory_date: date | None
    inventory_date_label: str
    revision: int
    generated_at_label: str
    total_records: int
    total_pieces: int
    summary: list[tuple[str, object]]
    observation: str
    records: list[dict[str, Any]]
    bay_groups: list[dict[str, Any]]
    consolidated_rows: list[list[object]]
    consolidated_kinds: list[str]
    conference_rows: list[list[object]]
    conference_kinds: list[str]
    has_reference: bool
    conciliation_rows: list[list[object]] = field(default_factory=list)
    conciliation_kinds: list[str] = field(default_factory=list)

    @property
    def subtitle(self) -> str:
        return f"Data do inventário: {self.inventory_date_label} · Revisão: {self.revision}"

    @property
    def generated_line(self) -> str:
        return f"{self.subtitle} · Gerado em: {self.generated_at_label}"

    def inventory_rows(self) -> list[list[object]]:
        return [
            [record["side"], record["bay"], record.get("layer") or "—", str(record["lot"]), int(record["quantity"])]
            for record in self.records
        ] + [["TOTAL", "", "", "", self.total_pieces]]

    def inventory_kinds(self) -> list[str]:
        return [str(record["side"]) for record in self.records] + ["total"]


def _summary(report: Mapping[str, Any], records: list[Mapping[str, Any]]) -> tuple[list[tuple[str, object]], str]:
    summary = report["summary"]
    side_pieces = {"DE": 0, "EF": 0}
    side_counts = {"DE": 0, "EF": 0}
    for record in records:
        side = str(record["side"])
        if side in side_pieces:
            side_pieces[side] += int(record["quantity"])
            side_counts[side] += 1
    lots_for_conference = int(summary.get("lotsForConference", summary.get("fragmentedLots", 0)))
    rows: list[tuple[str, object]] = [
        ("Total de registros", int(report["totalRecords"])),
        ("Total de peças", int(report["totalPieces"])),
        ("Total de lotes", int(summary["lotsAnalyzed"])),
        ("Total de vãos", len(_bay_groups(records))),
        ("Lotes OK", int(summary.get("lotsOk", summary.get("regularLots", 0)))),
        ("Lotes para conferência", lots_for_conference),
        ("1 peça fora do local principal", int(summary.get("singlePieceOutsideLots", summary.get("loosePieces", 0)))),
        ("Lotes com múltiplas peças fora do local principal", int(summary.get("multiplePiecesOutsideLots", summary.get("displacedGroups", 0)))),
        ("Lotes distribuídos em mais de um local", int(summary.get("distributedLots", summary.get("ambiguousDistributions", 0)))),
        ("Peças lado DE", side_pieces["DE"]),
        ("Peças lado EF", side_pieces["EF"]),
        ("Registros lado DE", side_counts["DE"]),
        ("Registros lado EF", side_counts["EF"]),
    ]
    if _has_reference(report):
        reference = report["reference"]
        rows.extend(
            [
                ("Referência de lotes", "Planilha Excel exportada do SAP / SICLA"),
                ("Lotes previstos", int(reference.get("totalLots", 0))),
                ("Previstos e encontrados", int(reference.get("foundLots", 0))),
                ("Previstos e não encontrados", int(reference.get("missingLots", 0))),
                ("Lotes físicos fora da referência", int(reference.get("outsideReferenceLots", 0))),
                ("Lotes fragmentados (físico)", int(reference.get("fragmentedLots", 0))),
            ]
        )
    if lots_for_conference:
        observation = (
            f"Há {format_int(lots_for_conference)} lote(s) para conferência. "
            "Consulte a seção LOTES PARA CONFERÊNCIA para verificar locais e quantidades."
        )
    else:
        observation = "Nenhum lote precisa de conferência."
    return rows, observation


def _consolidated(report: Mapping[str, Any]) -> tuple[list[list[object]], list[str]]:
    rows: list[list[object]] = []
    kinds: list[str] = []
    for lot in report["lots"]:
        presentation = _presentation(lot)
        requires_conference = bool(presentation["requiresConference"])
        location_summary = "\n".join(
            str(location["display"] if requires_conference else location["label"])
            for location in presentation["locations"]
        ) or "—"
        rows.append([str(lot["lot"]), int(lot["totalQuantity"]), location_summary, str(presentation["situation"])])
        kinds.append(str(presentation.get("tone", "review")))
    return rows, kinds


def _conference(report: Mapping[str, Any]) -> tuple[list[list[object]], list[str]]:
    """Uma linha por lote que exige conferência; Peças fora é None quando não se aplica."""

    rows: list[list[object]] = []
    kinds: list[str] = []
    for lot in report["lots"]:
        presentation = _presentation(lot)
        if not presentation["requiresConference"]:
            continue
        primary = presentation.get("primaryLocation")
        out_of_primary = presentation.get("outOfPrimaryQuantity")
        rows.append(
            [
                str(lot["lot"]),
                int(lot["totalQuantity"]),
                str(presentation["situation"]),
                str(primary["display"] if primary else "Não definido"),
                "\n".join(str(location["display"]) for location in presentation.get("otherLocations", [])) or "—",
                int(out_of_primary) if out_of_primary is not None else None,
                str(presentation["action"]),
            ]
        )
        kinds.append(str(presentation["tone"]))
    if not rows:
        return [list(NO_CONFERENCE_ROW)], ["ok"]
    return rows, kinds


def _reference_label(lot: Mapping[str, Any]) -> str:
    status = lot.get("referenceStatus")
    if status in {"EXPECTED_FOUND", "EXPECTED_MISSING"}:
        return "Previsto"
    if status == "OUTSIDE_REFERENCE":
        return "Fora da referência"
    return "Sem referência"


def _physical_label(lot: Mapping[str, Any]) -> str:
    return "Encontrado" if lot.get("physicalFound", int(lot.get("totalQuantity", 0)) > 0) else "Não encontrado"


def _conciliation(report: Mapping[str, Any]) -> tuple[list[list[object]], list[str]]:
    rows: list[list[object]] = []
    kinds: list[str] = []
    for lot in report.get("lots", []):
        if not isinstance(lot, Mapping):
            continue
        presentation = _presentation(lot)
        locations = presentation.get("locations", []) if isinstance(presentation, Mapping) else []
        location_summary = "\n".join(
            str(location.get("display") or location.get("label"))
            for location in locations
            if isinstance(location, Mapping)
        ) or "—"
        condition = str(presentation.get("situation", "LOTE PARA CONFERÊNCIA"))
        if lot.get("referenceStatus") == "EXPECTED_MISSING":
            condition = "PREVISTO E NÃO ENCONTRADO"
        elif lot.get("referenceStatus") == "OUTSIDE_REFERENCE":
            condition = f"FORA DA REFERÊNCIA · {condition}"
        rows.append(
            [str(lot["lot"]), _reference_label(lot), _physical_label(lot), location_summary, int(lot.get("totalQuantity", 0)), condition]
        )
        if lot.get("referenceStatus") in {"EXPECTED_MISSING", "OUTSIDE_REFERENCE"}:
            kinds.append("review")
        else:
            kinds.append(str(presentation.get("tone", "review")))
    return rows, kinds


def build_report_view(report: Mapping[str, Any]) -> ReportView:
    """Valida os totais e prepara a apresentação usada por todos os formatos."""

    verify_report_totals(report)
    records = [dict(record) for record in report["records"]]
    summary, observation = _summary(report, records)
    consolidated_rows, consolidated_kinds = _consolidated(report)
    conference_rows, conference_kinds = _conference(report)
    has_reference = _has_reference(report)
    conciliation_rows, conciliation_kinds = _conciliation(report) if has_reference else ([], [])
    inventory_date = parse_report_date(report["inventoryDate"])
    return ReportView(
        inventory_date=inventory_date,
        inventory_date_label=format_report_date(report["inventoryDate"]),
        revision=int(report["revision"]),
        generated_at_label=format_generated_at(report.get("generatedAt")),
        total_records=int(report["totalRecords"]),
        total_pieces=int(report["totalPieces"]),
        summary=summary,
        observation=observation,
        records=records,
        bay_groups=_bay_groups(records),
        consolidated_rows=consolidated_rows,
        consolidated_kinds=consolidated_kinds,
        conference_rows=conference_rows,
        conference_kinds=conference_kinds,
        has_reference=has_reference,
        conciliation_rows=conciliation_rows,
        conciliation_kinds=conciliation_kinds,
    )


def _excel_summary_rows(view: ReportView) -> list[list[object]]:
    """No Excel a data fica como célula de data e a observação vira a última linha."""

    first: list[list[object]] = [["Data do inventário", view.inventory_date or view.inventory_date_label]]
    return first + [[label, value] for label, value in view.summary] + [["Observações importantes", view.observation]]


def _excel_summary_subtitle(view: ReportView) -> str:
    return f"Revisão: {view.revision} · Gerado em: {view.generated_at_label}"


def _excel_inventory_rows(view: ReportView) -> list[list[object]]:
    rows = view.inventory_rows()
    for row in rows[:-1]:
        row[1] = _bay_value(str(row[1]))
    return rows


EXCEL_FOOTER = "Página &P de &N"
EXCEL_DATE_FORMAT = "DD/MM/YYYY"
EXCEL_THOUSANDS_FORMAT = "#,##0"


def _xlsx_add_sheet(
    workbook: Workbook,
    name: str,
    headers: list[str],
    rows: list[list[object]],
    *,
    title: str,
    subtitle: str,
    widths: list[int],
    text_columns: set[int] | None = None,
    number_columns: set[int] | None = None,
    row_kinds: list[str] | None = None,
    landscape: bool = True,
) -> None:
    sheet = workbook.create_sheet(name)
    last_column = len(headers)
    last_letter = get_column_letter(last_column)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(1, 1, title)
    sheet.cell(2, 1, subtitle)
    sheet.cell(1, 1).font = Font(name="Arial", size=15, bold=True, color=WHITE)
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=BLUE_DARK)
    sheet.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    sheet.cell(2, 1).font = Font(name="Arial", size=10, italic=True, color=BLUE_DARK)
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    sheet.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 20
    sheet.row_dimensions[3].height = 8

    header_row = 4
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column_index, header)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 30

    text_columns = text_columns or set()
    number_columns = number_columns or set()
    row_kinds = row_kinds or ["body"] * len(rows)
    border = Border(
        left=Side(style="thin", color=GREY_BORDER),
        right=Side(style="thin", color=GREY_BORDER),
        top=Side(style="thin", color=GREY_BORDER),
        bottom=Side(style="thin", color=GREY_BORDER),
    )
    white_fill = PatternFill("solid", fgColor=WHITE)
    alternate_fill = PatternFill("solid", fgColor=BLUE_PALE)
    for row_offset, values in enumerate(rows):
        row_number = header_row + 1 + row_offset
        kind = row_kinds[row_offset] if row_offset < len(row_kinds) else "body"
        row_fill = PatternFill("solid", fgColor=_ALERT_COLORS[kind]) if kind in _ALERT_COLORS else (alternate_fill if row_offset % 2 else white_fill)
        row_font = Font(name="Arial", size=10, color="000000")
        if kind == "total":
            row_fill = PatternFill("solid", fgColor=BLUE_DARK)
            row_font = Font(name="Arial", size=10, bold=True, color=WHITE)
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column_index, value)
            cell.font = row_font
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if isinstance(value, date):
                cell.number_format = EXCEL_DATE_FORMAT
            elif column_index in text_columns:
                cell.number_format = "@"
            elif column_index in number_columns:
                cell.number_format = EXCEL_THOUSANDS_FORMAT
            elif isinstance(value, int):
                cell.number_format = "0"
        if name == "INVENTÁRIO" and values and values[0] in {"DE", "EF"}:
            side_cell = sheet.cell(row_number, 1)
            side_cell.font = Font(name="Arial", size=10, bold=True, color=BLUE_DARK)
            side_cell.fill = PatternFill("solid", fgColor=BLUE_LIGHT)
            side_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row_number].height = 26 if any(len(str(value or "")) > 34 for value in values) else 20

    last_row = max(header_row, header_row + len(rows))
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False
    sheet.print_options.gridLines = False
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    sheet.oddFooter.center.text = EXCEL_FOOTER
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:{last_letter}{last_row}"


def generate_xlsx_report(report: dict[str, Any]) -> bytes:
    """Gera o relatório moderno em XLSX usando valores calculados no backend."""

    view = build_report_view(report)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _xlsx_add_sheet(
        workbook,
        "RESUMO",
        ["Indicador", "Valor"],
        _excel_summary_rows(view),
        title="Aplicativo Inventário da Laminação de Perfis",
        subtitle=_excel_summary_subtitle(view),
        widths=[38, 92],
        number_columns={2},
        landscape=False,
    )
    _xlsx_add_sheet(
        workbook,
        "INVENTÁRIO",
        INVENTORY_HEADERS,
        _excel_inventory_rows(view),
        title="INVENTÁRIO",
        subtitle=view.subtitle,
        widths=[12, 12, 14, 18, 22],
        text_columns={4},
        number_columns={5},
        row_kinds=view.inventory_kinds(),
    )
    _xlsx_add_sheet(
        workbook,
        "LOTES CONSOLIDADOS",
        CONSOLIDATED_HEADERS,
        view.consolidated_rows,
        title="LOTES CONSOLIDADOS",
        subtitle=view.subtitle,
        widths=[18, 18, 70, 46],
        text_columns={1},
        number_columns={2},
        row_kinds=view.consolidated_kinds,
    )
    _xlsx_add_sheet(
        workbook,
        "LOTES PARA CONFERÊNCIA",
        CONFERENCE_HEADERS,
        view.conference_rows,
        title="LOTES PARA CONFERÊNCIA",
        subtitle=view.subtitle,
        widths=[18, 16, 42, 34, 46, 16, 72],
        text_columns={1},
        number_columns={2, 6},
        row_kinds=view.conference_kinds,
    )
    if view.has_reference:
        _xlsx_add_sheet(
            workbook,
            "CONCILIAÇÃO",
            CONCILIATION_HEADERS,
            view.conciliation_rows,
            title="CONCILIAÇÃO COM REFERÊNCIA DE LOTES",
            subtitle=view.subtitle,
            widths=[18, 24, 18, 70, 16, 54],
            text_columns={1},
            number_columns={5},
            row_kinds=view.conciliation_kinds,
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_xlsx(report: dict[str, Any]) -> bytes:
    """Compatibilidade para o nome público usado pelo backend atual."""

    return generate_xlsx_report(report)


_XLS_BORDER = "borders: left thin, right thin, top thin, bottom thin;"
_XLS_STYLE_SPECS = {
    "title": "font: name Arial, height 280, bold on, colour white; pattern: pattern solid, fore_colour dark_blue; align: horiz left, vert centre;",
    "subtitle": "font: name Arial, height 200, italic on, colour dark_blue; pattern: pattern solid, fore_colour pale_blue; align: horiz left, vert centre;",
    "header": "font: name Arial, height 200, bold on, colour white; pattern: pattern solid, fore_colour dark_blue; align: horiz center, vert centre, wrap on;",
    "body": "font: name Arial, height 200; align: horiz center, vert centre, wrap on;",
    "body_alt": "font: name Arial, height 200; pattern: pattern solid, fore_colour pale_blue; align: horiz center, vert centre, wrap on;",
    "body_de": "font: name Arial, height 200, bold on, colour dark_blue; pattern: pattern solid, fore_colour pale_blue; align: horiz center, vert centre;",
    "body_ef": "font: name Arial, height 200, bold on, colour dark_blue; pattern: pattern solid, fore_colour ice_blue; align: horiz center, vert centre;",
    "total": "font: name Arial, height 200, bold on, colour white; pattern: pattern solid, fore_colour dark_blue; align: horiz center, vert centre;",
    "alert_red": "font: name Arial, height 200; pattern: pattern solid, fore_colour rose; align: horiz center, vert centre, wrap on;",
    "alert_orange": "font: name Arial, height 200; pattern: pattern solid, fore_colour light_orange; align: horiz center, vert centre, wrap on;",
    "alert_yellow": "font: name Arial, height 200; pattern: pattern solid, fore_colour light_yellow; align: horiz center, vert centre, wrap on;",
    "good": "font: name Arial, height 200; pattern: pattern solid, fore_colour light_green; align: horiz center, vert centre, wrap on;",
}


class _XlsStyles:
    """Cache de estilos BIFF8 por nome e formato numérico."""

    def __init__(self) -> None:
        self._cache: dict[tuple[str, str], xlwt.XFStyle] = {}

    def get(self, name: str, number_format: str = "General") -> xlwt.XFStyle:
        key = (name, number_format)
        if key not in self._cache:
            self._cache[key] = xlwt.easyxf(f"{_XLS_STYLE_SPECS[name]} {_XLS_BORDER}", num_format_str=number_format)
        return self._cache[key]


def _xls_row_style_name(kind: str, row_offset: int) -> str:
    if kind == "total":
        return "total"
    if kind == "single-piece":
        return "alert_yellow"
    if kind in {"multiple-pieces", "review"}:
        return "alert_orange"
    if kind == "distributed":
        return "alert_red"
    if kind == "ok":
        return "good"
    return "body_alt" if row_offset % 2 else "body"


def _xls_row_height(values: list[object]) -> int:
    """Reserva altura BIFF8 suficiente para texto quebrado e múltiplos locais."""

    line_count = max(
        (str(value or "").count("\n") + 1 for value in values),
        default=1,
    )
    if any(len(str(value or "")) > 34 for value in values):
        line_count = max(line_count, 2)
    return min(1200, 300 * line_count)


def _xls_add_sheet(
    workbook: xlwt.Workbook,
    styles: _XlsStyles,
    name: str,
    headers: list[str],
    rows: list[list[object]],
    *,
    title: str,
    subtitle: str,
    widths: list[int],
    text_columns: set[int] | None = None,
    number_columns: set[int] | None = None,
    row_kinds: list[str] | None = None,
    landscape: bool = True,
    print_scaling: int = 85,
) -> None:
    sheet = workbook.add_sheet(name)
    text_columns = text_columns or set()
    number_columns = number_columns or set()
    row_kinds = row_kinds or ["body"] * len(rows)
    sheet.write_merge(0, 0, 0, len(headers) - 1, title, styles.get("title"))
    sheet.write_merge(1, 1, 0, len(headers) - 1, subtitle, styles.get("subtitle"))
    for column_index, header in enumerate(headers):
        sheet.write(3, column_index, header, styles.get("header"))
    sheet.row(0).height = 560
    sheet.row(1).height = 360
    sheet.row(3).height = 520
    for column_index, width in enumerate(widths):
        sheet.col(column_index).width = min(width * 256, 65535)

    for row_offset, values in enumerate(rows):
        row_number = 4 + row_offset
        kind = row_kinds[row_offset] if row_offset < len(row_kinds) else "body"
        base_name = _xls_row_style_name(kind, row_offset)
        for column_index, value in enumerate(values):
            style_name = base_name
            if name == "INVENTÁRIO" and column_index == 0 and value in {"DE", "EF"}:
                style_name = "body_de" if value == "DE" else "body_ef"
            if isinstance(value, date):
                number_format = EXCEL_DATE_FORMAT
            elif column_index in text_columns:
                number_format = "@"
                value = str(value)
            elif column_index in number_columns:
                number_format = EXCEL_THOUSANDS_FORMAT
            elif isinstance(value, int):
                number_format = "0"
            else:
                number_format = "General"
            sheet.write(row_number, column_index, value, styles.get(style_name, number_format))
        sheet.row(row_number).height = _xls_row_height(values)

    sheet.set_panes_frozen(True)
    sheet.set_horz_split_pos(4)
    sheet.set_vert_split_pos(0)
    sheet.show_grid = False
    sheet.print_grid = False
    sheet.set_portrait(not landscape)
    sheet.set_paper_size_code(9)  # A4 no catálogo BIFF8.
    sheet.set_fit_width_to_pages(1)
    sheet.set_fit_height_to_pages(0)
    sheet.set_print_scaling(print_scaling)
    # O setter público do xlwt tenta decodificar `str` no Python 3 e falha;
    # o campo interno recebe o texto já em Unicode, como os valores padrão.
    sheet._Worksheet__header_str = ""
    sheet._Worksheet__footer_str = f"&C{EXCEL_FOOTER}"


def generate_xls_report(report: dict[str, Any]) -> bytes:
    """Gera diretamente um workbook BIFF8/Excel 97-2003, sem conversão."""

    view = build_report_view(report)
    workbook = xlwt.Workbook(encoding="utf-8")
    styles = _XlsStyles()
    _xls_add_sheet(
        workbook,
        styles,
        "RESUMO",
        ["Indicador", "Valor"],
        _excel_summary_rows(view),
        title="Aplicativo Inventário da Laminação de Perfis",
        subtitle=_excel_summary_subtitle(view),
        widths=[38, 92],
        number_columns={1},
        landscape=False,
        print_scaling=60,
    )
    _xls_add_sheet(
        workbook,
        styles,
        "INVENTÁRIO",
        INVENTORY_HEADERS,
        _excel_inventory_rows(view),
        title="INVENTÁRIO",
        subtitle=view.subtitle,
        widths=[12, 12, 14, 18, 22],
        text_columns={3},
        number_columns={4},
        row_kinds=view.inventory_kinds(),
    )
    _xls_add_sheet(
        workbook,
        styles,
        "LOTES CONSOLIDADOS",
        CONSOLIDATED_HEADERS,
        view.consolidated_rows,
        title="LOTES CONSOLIDADOS",
        subtitle=view.subtitle,
        widths=[18, 18, 70, 46],
        text_columns={0},
        number_columns={1},
        row_kinds=view.consolidated_kinds,
    )
    _xls_add_sheet(
        workbook,
        styles,
        "LOTES PARA CONFERÊNCIA",
        CONFERENCE_HEADERS,
        view.conference_rows,
        title="LOTES PARA CONFERÊNCIA",
        subtitle=view.subtitle,
        widths=[18, 16, 42, 34, 46, 16, 72],
        text_columns={0},
        number_columns={1, 5},
        row_kinds=view.conference_kinds,
        print_scaling=60,
    )
    if view.has_reference:
        _xls_add_sheet(
            workbook,
            styles,
            "CONCILIAÇÃO",
            CONCILIATION_HEADERS,
            view.conciliation_rows,
            title="CONCILIAÇÃO COM REFERÊNCIA DE LOTES",
            subtitle=view.subtitle,
            widths=[18, 24, 18, 70, 16, 54],
            text_columns={0},
            number_columns={4},
            row_kinds=view.conciliation_kinds,
            print_scaling=55,
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_xls(report: dict[str, Any]) -> bytes:
    """Compatibilidade para o nome público do gerador legado."""

    return generate_xls_report(report)


def _display(value: object) -> str:
    """Texto de célula para PDF e Word: inteiros com milhar e vazio como travessão."""

    if value is None or value == "":
        return "—"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return format_int(value)
    return str(value)


def _display_rows(rows: list[list[object]]) -> list[list[str]]:
    return [[_display(value) for value in row] for row in rows]


def _summary_display_rows(view: ReportView) -> list[list[str]]:
    return [[label, _display(value)] for label, value in view.summary]


def _inventory_display_rows(view: ReportView) -> list[list[str]]:
    rows = _display_rows(view.inventory_rows())
    rows[-1][1:4] = ["", "", ""]
    return rows


def _pdf_value(value: object) -> str:
    return escape(str(value)).replace("\n", "<br/>")


def _pdf_table(rows: list[list[object]], widths: list[float], row_kinds: list[str] | None = None) -> Table:
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "report-cell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=TA_LEFT,
        spaceAfter=0,
        splitLongWords=0,
    )
    header_style = ParagraphStyle("report-header", parent=cell_style, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_LEFT)
    total_style = ParagraphStyle("report-total", parent=cell_style, fontName="Helvetica-Bold", textColor=colors.white)
    kinds = ["header"] + list(row_kinds or [])
    rendered = [
        [
            Paragraph(
                _pdf_value(value),
                header_style if row_index == 0 else total_style if row_index < len(kinds) and kinds[row_index] == "total" else cell_style,
            )
            for value in row
        ]
        for row_index, row in enumerate(rows)
    ]
    table = Table(rendered, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands: list[tuple[Any, ...]] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE_DARK}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor(f"#{GREY_BORDER}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{BLUE_PALE}")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for row_offset, kind in enumerate(row_kinds or [], start=1):
        if kind == "total":
            commands.append(("BACKGROUND", (0, row_offset), (-1, row_offset), colors.HexColor(f"#{BLUE_DARK}")))
            commands.append(("TEXTCOLOR", (0, row_offset), (-1, row_offset), colors.white))
        elif kind in _ALERT_COLORS:
            commands.append(("BACKGROUND", (0, row_offset), (-1, row_offset), colors.HexColor(f"#{_ALERT_COLORS[kind]}")))
    table.setStyle(TableStyle(commands))
    return table


def _pdf_page_frame(canvas: Any, document: Any) -> None:
    page_width, page_height = document.pagesize
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor(f"#{GREY_BORDER}"))
    canvas.line(document.leftMargin, page_height - 1.35 * cm, page_width - document.rightMargin, page_height - 1.35 * cm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor(f"#{BLUE_DARK}"))
    canvas.drawString(document.leftMargin, page_height - 1.05 * cm, "INVENTÁRIO · RELATÓRIO OFICIAL")
    canvas.restoreState()


class _NumberedCanvas(pdf_canvas.Canvas):
    """Adia o rodapé até conhecer o total de páginas (Página X de Y)."""

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._page_states: list[dict[str, Any]] = []

    def showPage(self) -> None:  # noqa: N802 - API do ReportLab
        self._page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        total = len(self._page_states)
        for state in self._page_states:
            self.__dict__.update(state)
            page_width = self._pagesize[0]
            self.saveState()
            self.setFont("Helvetica", 7)
            self.setFillColor(colors.HexColor(f"#{BLUE_DARK}"))
            self.drawRightString(page_width - 1.3 * cm, 0.75 * cm, f"Página {self._pageNumber} de {total}")
            self.restoreState()
            super().showPage()
        super().save()


# Espaço mínimo para um título seguido do cabeçalho e das primeiras linhas da tabela.
_PDF_HEADING_MIN_SPACE = 3.2 * cm


def export_pdf(report: dict[str, Any]) -> bytes:
    view = build_report_view(report)
    buffer = BytesIO()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("report-title", parent=styles["Title"], textColor=colors.HexColor(f"#{BLUE_DARK}"), fontSize=18, leading=22, alignment=TA_LEFT, spaceAfter=4)
    subtitle_style = ParagraphStyle("report-subtitle", parent=styles["Normal"], textColor=colors.HexColor("#4F5B66"), fontSize=9, leading=12, spaceAfter=8)
    heading_style = ParagraphStyle("report-heading", parent=styles["Heading2"], textColor=colors.HexColor(f"#{BLUE_DARK}"), fontSize=12, leading=14, spaceBefore=8, spaceAfter=6, keepWithNext=1)
    note_style = ParagraphStyle("report-note", parent=styles["BodyText"], fontSize=9, leading=12, spaceBefore=5, spaceAfter=5)
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.3 * cm,
        rightMargin=1.3 * cm,
        topMargin=1.7 * cm,
        bottomMargin=1.3 * cm,
        title="Relatório de Inventário",
    )

    def heading(text: str) -> list[Any]:
        return [CondPageBreak(_PDF_HEADING_MIN_SPACE), Paragraph(text, heading_style)]

    story: list[Any] = [
        Paragraph("Aplicativo Inventário da Laminação de Perfis", title_style),
        Paragraph("Relatório de Inventário", heading_style),
        Paragraph(escape(view.generated_line), subtitle_style),
        *heading("Resumo geral"),
        _pdf_table([["Indicador", "Valor"]] + _summary_display_rows(view), [14.5 * cm, 11.5 * cm]),
        *heading("Observação operacional"),
        Paragraph(escape(view.observation), note_style),
        Spacer(1, 0.15 * cm),
        *heading("Inventário"),
        _pdf_table(
            [INVENTORY_HEADERS] + _inventory_display_rows(view),
            [2.0 * cm, 2.0 * cm, 2.4 * cm, 15.0 * cm, 4.0 * cm],
            view.inventory_kinds(),
        ),
        PageBreak(),
        Paragraph("Lotes consolidados", heading_style),
        _pdf_table(
            [CONSOLIDATED_HEADERS] + _display_rows(view.consolidated_rows),
            [2.0 * cm, 2.5 * cm, 13.3 * cm, 9.3 * cm],
            view.consolidated_kinds,
        ),
        PageBreak(),
        Paragraph("Lotes para conferência", heading_style),
        _pdf_table(
            [CONFERENCE_HEADERS] + _display_rows(view.conference_rows),
            [2.0 * cm, 1.6 * cm, 4.5 * cm, 4.2 * cm, 5.0 * cm, 1.7 * cm, 8.1 * cm],
            view.conference_kinds,
        ),
    ]
    if view.has_reference:
        story.extend(
            [
                PageBreak(),
                Paragraph("Conciliação com referência de lotes", heading_style),
                Paragraph("A referência usa somente números de lote da planilha Excel exportada do SAP / SICLA. As quantidades abaixo são exclusivamente físicas.", note_style),
                _pdf_table(
                    [CONCILIATION_HEADERS] + _display_rows(view.conciliation_rows),
                    [2.0 * cm, 3.0 * cm, 2.6 * cm, 10.2 * cm, 2.5 * cm, 7.0 * cm],
                    view.conciliation_kinds,
                ),
            ]
        )
    document.build(story, onFirstPage=_pdf_page_frame, onLaterPages=_pdf_page_frame, canvasmaker=_NumberedCanvas)
    return buffer.getvalue()


def _docx_set_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _docx_set_row_flags(row: Any, *, header: bool = False) -> None:
    """Repete o cabeçalho em cada página e impede linha partida entre páginas."""

    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    properties.append(cant_split)
    if header:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        properties.append(repeat)


def _docx_cell_text(cell: Any, value: object, *, header: bool = False, white: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    run.font.name = "Arial"
    run.font.size = Pt(8 if header else 7.5)
    run.font.bold = header or white
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)


def _docx_add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    row_kinds: list[str] | None = None,
) -> Any:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _docx_set_row_flags(table.rows[0], header=True)
    for column_index, header in enumerate(headers):
        cell = table.rows[0].cells[column_index]
        cell.width = Cm(widths[column_index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _docx_set_shading(cell, BLUE_DARK)
        _docx_cell_text(cell, header, header=True, white=True)
    for row_offset, values in enumerate(rows):
        kind = (row_kinds or ["body"] * len(rows))[row_offset]
        row = table.add_row()
        _docx_set_row_flags(row)
        fill = _ALERT_COLORS.get(kind, BLUE_PALE if row_offset % 2 else WHITE)
        white = kind == "total"
        if white:
            fill = BLUE_DARK
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.width = Cm(widths[column_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _docx_set_shading(cell, fill)
            _docx_cell_text(cell, value, white=white)
    return table


def _docx_add_field(paragraph: Any, instruction_text: str) -> None:
    run = paragraph.add_run()
    run.font.name = "Arial"
    run.font.size = Pt(8)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {instruction_text} "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def _docx_add_text(paragraph: Any, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(8)


def _docx_add_page_field(paragraph: Any) -> None:
    _docx_add_text(paragraph, "Página ")
    _docx_add_field(paragraph, "PAGE")
    _docx_add_text(paragraph, " de ")
    _docx_add_field(paragraph, "NUMPAGES")


def _docx_heading(document: Document, text: str) -> Any:
    heading = document.add_heading(text, level=1)
    heading.paragraph_format.keep_with_next = True
    return heading


def export_docx(report: dict[str, Any]) -> bytes:
    view = build_report_view(report)
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    header = section.header.paragraphs[0]
    header.text = "INVENTÁRIO · RELATÓRIO OFICIAL"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(31, 78, 120)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _docx_add_page_field(footer)

    document.core_properties.title = "Relatório de Inventário"
    title = document.add_heading("Aplicativo Inventário da Laminação de Perfis", level=0)
    title.runs[0].font.name = "Arial"
    title.runs[0].font.color.rgb = RGBColor(31, 78, 120)
    document.add_paragraph("Relatório de Inventário")
    document.add_paragraph(view.generated_line)
    _docx_heading(document, "Resumo geral")
    _docx_add_table(document, ["Indicador", "Valor"], _summary_display_rows(view), [7.0, 10.0])
    _docx_heading(document, "Observação operacional")
    document.add_paragraph(view.observation)
    _docx_heading(document, "Inventário")
    _docx_add_table(
        document,
        INVENTORY_HEADERS,
        _inventory_display_rows(view),
        [1.5, 1.45, 2.0, 7.0, 2.4],
        row_kinds=view.inventory_kinds(),
    )
    document.add_page_break()
    _docx_heading(document, "Lotes consolidados")
    _docx_add_table(
        document,
        CONSOLIDATED_HEADERS,
        _display_rows(view.consolidated_rows),
        [2.4, 2.5, 14.5, 7.5],
        row_kinds=view.consolidated_kinds,
    )
    document.add_page_break()
    _docx_heading(document, "Lotes para conferência")
    _docx_add_table(
        document,
        CONFERENCE_HEADERS,
        _display_rows(view.conference_rows),
        [2.2, 1.5, 4.5, 4.2, 5.0, 2.0, 8.0],
        row_kinds=view.conference_kinds,
    )
    if view.has_reference:
        document.add_page_break()
        _docx_heading(document, "Conciliação com referência de lotes")
        document.add_paragraph("A referência usa somente números de lote da planilha Excel exportada do SAP / SICLA. As quantidades abaixo são exclusivamente físicas.")
        _docx_add_table(
            document,
            CONCILIATION_HEADERS,
            _display_rows(view.conciliation_rows),
            [2.2, 3.0, 2.4, 10.6, 2.4, 7.0],
            row_kinds=view.conciliation_kinds,
        )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
