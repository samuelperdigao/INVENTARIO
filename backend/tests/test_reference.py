from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select
from xlrd import open_workbook

from app.database import get_session
from app.engine import AnalysisEntry
from app.main import app
from app.persistence import InventoryEntryRow, InventoryReferenceRow, ReferenceLotRow
from app.reference_service import ReferenceImportError, parse_xlsx_reference
from app.reports import build_consolidated_report, export_docx, export_pdf, export_xls, export_xlsx
from auth_helpers import register_verified


client = TestClient(app)


def _xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_with_number_format() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Descrição", "Lotes"])
    worksheet.append(["Texto", "2712345678"])
    worksheet.append(["Número", 2712345679])
    worksheet["B3"].number_format = "0000000000"
    worksheet.append(["Notação", "2,71234568E+9"])
    worksheet.append(["Duplicado", "2712345678\u200b"])
    worksheet.append(["Inválido", "abc"])
    worksheet.append(["Vazio", None])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _with_zip_member(content: bytes, name: str) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(content), "r") as source, zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr(name, b"invalid")
    return output.getvalue()


def _times() -> dict[str, str]:
    value = datetime.now(timezone.utc).isoformat()
    return {"createdAt": value, "updatedAt": value}


def _inventory_payload(inventory_id: str) -> dict[str, object]:
    return {
        "id": inventory_id,
        "date": "2026-09-16",
        "status": "OPEN",
        "revision": 1,
        "syncBaseRevision": 0,
        "tombstone": False,
        "deletedAt": None,
        **_times(),
    }


def _entry_payload(inventory_id: str, lot: str, side: str, bay: str, quantity: int) -> dict[str, object]:
    return {
        "id": str(uuid4()),
        "inventoryId": inventory_id,
        "side": side,
        "bay": bay,
        "lot": lot,
        "quantity": quantity,
        "revision": 1,
        "syncBaseRevision": 0,
        "tombstone": False,
        "deletedAt": None,
        **_times(),
    }


def _create_inventory(
    email: str,
    entries: list[dict[str, object]] | None = None,
) -> tuple[str, str, dict[str, str]]:
    _, auth = register_verified(client, email)
    inventory_id = str(uuid4())
    sync_token = f"sync-token-{uuid4()}-{uuid4()}"
    prepared_entries = [{**entry, "inventoryId": inventory_id} for entry in entries or []]
    response = client.post(
        "/api/v1/sync",
        json={
            "deviceId": str(uuid4()),
            "inventoryId": inventory_id,
            "teamId": None,
            "cursor": 0,
            "inventory": _inventory_payload(inventory_id),
            "entries": prepared_entries,
        },
        headers={**auth, "X-Inventory-Sync-Token": sync_token},
    )
    assert response.status_code == 200, response.text
    return inventory_id, sync_token, auth


def _upload(content: bytes, filename: str = "referencia.xlsx") -> dict[str, tuple[str, bytes, str]]:
    return {"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_parser_normalizes_text_numbers_scientific_values_and_duplicates() -> None:
    parsed = parse_xlsx_reference(
        _xlsx_with_number_format(),
        r"C:\importacoes\sap\referencia.xlsx",
        max_bytes=2 * 1024 * 1024,
        max_rows=100,
    )

    assert parsed.original_filename == "referencia.xlsx"
    assert parsed.selected_column == 2
    assert parsed.selected_column_label == "Lotes"
    assert parsed.total_rows == 6
    assert parsed.valid_lot_occurrences == 4
    assert parsed.unique_lots == 3
    assert parsed.duplicate_rows == 1
    assert parsed.ignored_rows == 2
    assert parsed.lots == ("2712345678", "2712345679", "2712345680")
    assert any("caracteres" in warning for warning in parsed.warnings)
    assert any("repetida" in warning for warning in parsed.warnings)


def test_parser_requires_exact_lotes_header_and_never_accepts_another_column() -> None:
    content = _xlsx([["Lotes", "Peso"], ["2712345678", "2812345678"]])

    preview = parse_xlsx_reference(content, "sap.xlsx", max_bytes=2 * 1024 * 1024, max_rows=100)
    assert preview.requires_column_selection is False
    assert preview.selected_column == 1
    assert preview.selected_column_label == "Lotes"
    assert preview.lots == ("2712345678",)

    with pytest.raises(ReferenceImportError, match="localizar a coluna"):
        parse_xlsx_reference(
            _xlsx([["Lote", "Peso"], ["2712345678", "2812345678"]]),
            "sap.xlsx",
            max_bytes=2 * 1024 * 1024,
            max_rows=100,
        )
    with pytest.raises(ReferenceImportError, match="exclusivamente"):
        parse_xlsx_reference(
            content,
            "sap.xlsx",
            selected_column=2,
            max_bytes=2 * 1024 * 1024,
            max_rows=100,
        )


def test_parser_accepts_case_and_outer_spacing_only_for_lotes_header() -> None:
    parsed = parse_xlsx_reference(
        _xlsx([["  LOTES  ", "Lotes SAP"], [2812345678, 2898765432]]),
        "sap.xlsx",
        max_bytes=2 * 1024 * 1024,
        max_rows=100,
    )

    assert parsed.selected_column == 1
    assert parsed.lots == ("2812345678",)


def test_parser_ignores_invalid_empty_and_other_column_values_without_false_positives() -> None:
    parsed = parse_xlsx_reference(
        _xlsx([
            ["Peso", "Lotes", "Material", "Descrição"],
            [2712345678, 2712345678, 2812345678, "lote 2898765432"],
            [2812345679, "2612345678", 2712345680, ""],
            [2812345681, None, 2812345682, ""],
        ]),
        "sap.xlsx",
        max_bytes=2 * 1024 * 1024,
        max_rows=100,
    )

    assert parsed.lots == ("2712345678",)
    assert parsed.valid_lot_occurrences == 1
    assert parsed.duplicate_rows == 0
    assert parsed.ignored_rows == 2
    assert any("10 números" in warning for warning in parsed.warnings)


def test_parser_returns_preview_with_warnings_when_lotes_has_no_valid_values() -> None:
    parsed = parse_xlsx_reference(
        _xlsx([["Lotes"], ["2612345678"], [None]]),
        "sap.xlsx",
        max_bytes=2 * 1024 * 1024,
        max_rows=100,
    )

    assert parsed.lots == ()
    assert parsed.unique_lots == 0
    assert any("Nenhum lote válido" in warning for warning in parsed.warnings)


def test_parser_rejects_zip_path_traversal() -> None:
    with pytest.raises(ReferenceImportError, match="caminho interno inválido"):
        parse_xlsx_reference(
            _with_zip_member(_xlsx([["Lotes"], ["2712345678"]]), "../escape.txt"),
            "sap.xlsx",
            max_bytes=2 * 1024 * 1024,
            max_rows=100,
        )


def test_reference_preview_import_replace_remove_and_match() -> None:
    inventory_id, sync_token, auth = _create_inventory(
        "referencia.owner@example.com",
        [
            _entry_payload("ignored", "2712345678", "DE", "15", 19),
            _entry_payload("ignored", "2898765432", "EF", "21", 1),
        ],
    )

    headers = {**auth, "X-Inventory-Sync-Token": sync_token}
    initial = _xlsx([["Lotes"], ["2712345678"], ["2812345678"]])
    preview = client.post(
        f"/api/v1/inventories/{inventory_id}/reference/preview",
        files=_upload(initial),
        headers=headers,
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["selectedColumn"] == 1
    assert preview.json()["uniqueLots"] == 2

    imported = client.post(
        f"/api/v1/inventories/{inventory_id}/reference",
        files=_upload(initial),
        data={"columnIndex": "1"},
        headers=headers,
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["lotNumbers"] == ["2712345678", "2812345678"]
    assert imported.json()["reference"]["totalLots"] == 2

    state = client.get(f"/api/v1/inventories/{inventory_id}/reference", headers=headers)
    assert state.status_code == 200
    assert state.json()["summary"] == {
        "available": True,
        "totalLots": 2,
        "foundLots": 1,
        "pendingLots": 1,
        "outsideReferenceLots": 1,
        "fragmentedLots": 0,
        "physicalDistinctLots": 2,
    }
    assert state.json()["lots"] == [
        {"lotNumber": "2712345678", "foundPhysically": True, "physicalQuantity": 19, "physicalOccurrences": 1, "fragmented": False},
        {"lotNumber": "2812345678", "foundPhysically": False, "physicalQuantity": 0, "physicalOccurrences": 0, "fragmented": False},
    ]
    assert client.get(f"/api/v1/inventories/{inventory_id}/reference/match/2712345678", headers=headers).json()["inReference"] is True
    assert client.get(f"/api/v1/inventories/{inventory_id}/reference/match/2898765432", headers=headers).json()["inReference"] is False

    replacement = _xlsx([["Lotes"], ["2898765432"]])
    replaced = client.post(
        f"/api/v1/inventories/{inventory_id}/reference",
        files=_upload(replacement, "substituicao.xlsx"),
        data={"columnIndex": "1"},
        headers=headers,
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["reference"]["revision"] == 2
    assert replaced.json()["lotNumbers"] == ["2898765432"]

    replaced_state = client.get(f"/api/v1/inventories/{inventory_id}/reference", headers=headers).json()
    assert replaced_state["summary"]["foundLots"] == 1
    assert replaced_state["summary"]["outsideReferenceLots"] == 1
    assert [item["lotNumber"] for item in replaced_state["lots"]] == ["2898765432"]

    removed = client.delete(f"/api/v1/inventories/{inventory_id}/reference", headers=headers)
    assert removed.status_code == 200
    assert client.get(f"/api/v1/inventories/{inventory_id}/reference", headers=headers).json()["reference"] is None
    assert client.get(f"/api/v1/inventories/{inventory_id}/reference/match/2898765432", headers=headers).json() == {
        "referenceAvailable": False,
        "lot": "2898765432",
        "inReference": None,
    }

    dependency = app.dependency_overrides[get_session]()
    session = next(dependency)
    try:
        assert session.scalar(select(func.count()).select_from(InventoryEntryRow).where(InventoryEntryRow.inventory_id == inventory_id)) == 2
        reference = session.scalar(select(InventoryReferenceRow).where(InventoryReferenceRow.inventory_id == inventory_id))
        assert reference is not None
        assert reference.status == "REMOVED"
        assert session.scalar(select(func.count()).select_from(ReferenceLotRow).where(ReferenceLotRow.reference_id == reference.id)) == 0
    finally:
        dependency.close()


def test_reference_rejects_invalid_file_without_touching_physical_inventory() -> None:
    inventory_id, sync_token, auth = _create_inventory(
        "referencia.invalid-file@example.com",
        [_entry_payload("ignored", "2712345678", "DE", "15", 19)],
    )
    headers = {**auth, "X-Inventory-Sync-Token": sync_token}

    missing_header = client.post(
        f"/api/v1/inventories/{inventory_id}/reference",
        files=_upload(_xlsx([["Lote"], ["2712345678"]]), "sem-cabecalho.xlsx"),
        headers=headers,
    )
    assert missing_header.status_code == 422
    assert "localizar a coluna" in missing_header.json()["detail"]

    invalid_rows = client.post(
        f"/api/v1/inventories/{inventory_id}/reference",
        files=_upload(_xlsx([["Lotes"], ["2612345678"], ["abc"]]), "sem-lotes-validos.xlsx"),
        headers=headers,
    )
    assert invalid_rows.status_code == 422
    assert "Nenhum lote válido" in invalid_rows.json()["detail"]

    state = client.get(f"/api/v1/inventories/{inventory_id}/reference", headers=headers)
    assert state.status_code == 200
    assert state.json()["reference"] is None
    physical = client.get(f"/api/v1/inventories/{inventory_id}/lots/2712345678", headers=headers)
    assert physical.status_code == 200
    assert len(physical.json()) == 1


def test_reference_endpoints_do_not_allow_idor() -> None:
    inventory_id, sync_token, owner_auth = _create_inventory("reference.idor.owner@example.com")
    _, attacker_auth = register_verified(client, "reference.idor.attacker@example.com")
    headers = {**attacker_auth, "X-Inventory-Sync-Token": sync_token}
    content = _xlsx([["Lotes"], ["2712345678"]])

    preview = client.post(
        f"/api/v1/inventories/{inventory_id}/reference/preview",
        files=_upload(content),
        headers=headers,
    )
    assert preview.status_code == 404
    assert client.get(f"/api/v1/inventories/{inventory_id}/reference", headers=attacker_auth).status_code == 404
    assert client.delete(f"/api/v1/inventories/{inventory_id}/reference", headers=headers).status_code == 404
    assert client.get(f"/api/v1/inventories/{inventory_id}/reference/match/2712345678", headers=attacker_auth).status_code == 404
    assert owner_auth["Authorization"] != attacker_auth["Authorization"]


def test_reference_conciliation_reaches_all_report_formats() -> None:
    report = build_consolidated_report(
        str(uuid4()),
        "2026-09-16",
        4,
        [
            # Previsto e encontrado.
            AnalysisEntry(side="DE", bay="15", lot="2712345678", quantity=19),
            # Físico fora da referência.
            AnalysisEntry(side="EF", bay="21", lot="2898765432", quantity=1),
        ],
        datetime(2026, 9, 16, tzinfo=timezone.utc),
        reference_lots=("2712345678", "2812345678"),
        reference_metadata={
            "sourceType": "SAP_EXCEL",
            "originalFilename": "sap.xlsx",
            "importedAt": "2026-09-16T10:00:00+00:00",
            "updatedAt": "2026-09-16T10:00:00+00:00",
            "revision": 1,
        },
    )

    lots = {lot["lot"]: lot for lot in report["lots"]}  # type: ignore[index]
    assert lots["2712345678"]["referenceStatus"] == "EXPECTED_FOUND"
    assert lots["2812345678"]["referenceStatus"] == "EXPECTED_MISSING"
    assert lots["2898765432"]["referenceStatus"] == "OUTSIDE_REFERENCE"
    assert report["reference"]["missingLots"] == 1  # type: ignore[index]

    xlsx = load_workbook(BytesIO(export_xlsx(report)))
    assert "CONCILIAÇÃO" in xlsx.sheetnames
    xlsx_rows = [xlsx["CONCILIAÇÃO"].cell(row, 1).value for row in range(5, xlsx["CONCILIAÇÃO"].max_row + 1)]
    assert xlsx_rows == ["2712345678", "2812345678", "2898765432"]

    xls = open_workbook(file_contents=export_xls(report))
    conciliation = xls.sheet_by_name("CONCILIAÇÃO")
    assert [conciliation.cell_value(row, 0) for row in range(4, conciliation.nrows)] == ["2712345678", "2812345678", "2898765432"]

    pdf_text = "".join(page.extract_text() or "" for page in PdfReader(BytesIO(export_pdf(report))).pages)
    assert "Conciliação com referência de lotes" in pdf_text
    assert "PREVISTO E NÃO ENCONTRADO" in pdf_text
    assert "FORA DA REFERÊNCIA" in pdf_text

    document = Document(BytesIO(export_docx(report)))
    docx_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    assert "Conciliação com referência de lotes" in docx_text
    assert "PREVISTO E NÃO ENCONTRADO" in docx_text
    assert "FORA DA REFERÊNCIA" in docx_text
