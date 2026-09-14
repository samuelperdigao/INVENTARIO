from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from xlrd import XL_CELL_TEXT, open_workbook

from app.engine import AnalysisEntry
from app.reports import build_consolidated_report, export_docx, export_pdf, export_xls, export_xlsx


def report() -> dict[str, object]:
    return build_consolidated_report(
        str(uuid4()), "2026-09-11", 3,
        [
            AnalysisEntry(side="DE", bay="15", layer="A1", lot="000123", quantity=19),
            AnalysisEntry(side="EF", bay="21", layer="A2", lot="000123", quantity=1),
            AnalysisEntry(side="DE", bay="08", lot="10", quantity=15),
            AnalysisEntry(side="EF", bay="11", lot="10", quantity=2),
            AnalysisEntry(side="DE", bay="20", lot="10", quantity=3),
            AnalysisEntry(side="EF", bay="01", lot="20", quantity=10),
            AnalysisEntry(side="DE", bay="01", lot="20", quantity=10),
            AnalysisEntry(side="EF", bay="02", lot="30", quantity=11),
            AnalysisEntry(side="DE", bay="02", lot="30", quantity=9),
        ],
        datetime(2026, 9, 11, tzinfo=timezone.utc),
    )


def test_consolidated_model_applies_required_classifications() -> None:
    lots = {lot["lot"]: lot for lot in report()["lots"]}  # type: ignore[index]
    assert lots["000123"]["classification"] == "PEÇA_SOLTEIRA"
    assert lots["10"]["classification"] == "GRUPO_DESLOCADO"
    assert lots["20"]["classification"] == "DISTRIBUIÇÃO_AMBÍGUA"
    assert lots["30"]["classification"] == "DISTRIBUIÇÃO_AMBÍGUA"


def test_xlsx_has_operational_tabs_layer_and_lot_as_text() -> None:
    workbook = load_workbook(BytesIO(export_xlsx(report())))
    assert workbook.sheetnames == ["RESUMO", "INVENTÁRIO", "LOTES CONSOLIDADOS", "DIVERGÊNCIAS"]
    inventory = workbook["INVENTÁRIO"]
    lot_cell = next(cell for cell in inventory["D"] if cell.value == "000123")
    assert lot_cell.number_format == "@"
    assert inventory.cell(lot_cell.row, 3).value in {"A1", "A2"}
    assert inventory.freeze_panes == "A5"
    assert inventory.auto_filter.ref is None
    assert inventory.sheet_view.showGridLines is False
    assert inventory.print_options.gridLines is False
    assert inventory["D5"].alignment.horizontal == "center"
    assert inventory.page_setup.paperSize == 9
    summary = {workbook["RESUMO"].cell(row, 1).value: workbook["RESUMO"].cell(row, 2).value for row in range(5, 22)}
    assert summary["Total de registros"] == 9
    assert summary["Total de lotes"] == 4
    assert summary["Total de peças"] == 80


def test_xls_is_biff8_with_same_tabs_and_text_lots() -> None:
    content = export_xls(report())
    assert len(content) > 0
    assert content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    workbook = open_workbook(file_contents=content)
    assert workbook.sheet_names() == ["RESUMO", "INVENTÁRIO", "LOTES CONSOLIDADOS", "DIVERGÊNCIAS"]
    assert sum(name.name == "_FilterDatabase" for name in workbook.name_obj_list) == 0
    assert all(not workbook.sheet_by_name(name).show_grid_lines for name in workbook.sheet_names())
    inventory = workbook.sheet_by_name("INVENTÁRIO")
    lot_row = next(row for row in range(inventory.nrows) if inventory.cell_value(row, 3) == "000123")
    assert inventory.cell_type(lot_row, 3) == XL_CELL_TEXT
    assert inventory.cell_value(lot_row, 2) in {"A1", "A2"}
    summary = {workbook.sheet_by_name("RESUMO").cell_value(row, 0): workbook.sheet_by_name("RESUMO").cell_value(row, 1) for row in range(4, 21)}
    assert summary["Total de registros"] == 9
    assert summary["Total de lotes"] == 4
    assert summary["Total de peças"] == 80


def test_pdf_and_docx_are_valid_and_contain_critical_content() -> None:
    pdf = PdfReader(BytesIO(export_pdf(report())))
    pdf_text = "".join(page.extract_text() or "" for page in pdf.pages)
    assert "Relatório de Inventário" in pdf_text
    assert "Camada" in pdf_text
    assert "Lotes consolidados" in pdf_text
    assert "Divergências e recomendações" in pdf_text
    document = Document(BytesIO(export_docx(report())))
    assert "Relatório de Inventário" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    raw_table = next(table for table in document.tables if len(table.rows[0].cells) >= 3 and table.rows[0].cells[2].text == "Camada")
    assert raw_table.rows[0].cells[3].text == "Lote"
    assert len(document.tables) == 4
