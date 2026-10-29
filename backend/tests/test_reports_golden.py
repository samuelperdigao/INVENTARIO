"""Testes golden com o inventário real de 06/10/2026 (revisão 6).

Cada formato de entrega é lido de volta: os lançamentos extraídos precisam ser
idênticos à entrada, os totais precisam fechar e os formatos precisam
concordar entre si e com o modelo de dados do relatório.
"""

import json
import re
from collections import Counter
from copy import deepcopy
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pypdf import PdfReader
from xlrd import XL_CELL_BLANK, XL_CELL_DATE, XL_CELL_NUMBER, XL_CELL_TEXT, open_workbook, xldate_as_datetime

from app.engine import AnalysisEntry
from app.reports import (
    ReportConsistencyError,
    build_inventory_report_data,
    build_report_view,
    export_docx,
    export_pdf,
    export_xls,
    export_xlsx,
    SECTION_TITLES,
    format_generated_at,
)


FIXTURE = json.loads((Path(__file__).parent / "fixtures" / "inventario_2026-10-06_rev6.json").read_text(encoding="utf-8"))
EXPECTED = FIXTURE["expected"]
LOT_PATTERN = re.compile(r"^2[78]\d{8}$")
Record = tuple[str, str, str | None, str, int]


def _input_records() -> list[Record]:
    return [(row["side"], row["bay"], row["layer"], row["lot"], row["quantity"]) for row in FIXTURE["records"]]


@pytest.fixture(scope="module")
def report() -> dict[str, object]:
    return build_inventory_report_data(
        "00000000-0000-4000-8000-000000000001",
        FIXTURE["inventoryDate"],
        FIXTURE["revision"],
        [AnalysisEntry(**row) for row in FIXTURE["records"]],
        datetime.fromisoformat(FIXTURE["generatedAt"]),
    )


@pytest.fixture(scope="module")
def files(report: dict[str, object]) -> dict[str, bytes]:
    return {"xls": export_xls(report), "xlsx": export_xlsx(report), "pdf": export_pdf(report), "docx": export_docx(report)}


def _layer(value: object) -> str | None:
    return None if value in (None, "", "—") else str(value)


def _int(value: object) -> int:
    return int(str(value).replace(".", "")) if isinstance(value, str) else int(value)  # type: ignore[arg-type]


def _bay(value: object) -> str:
    return str(int(value)) if isinstance(value, (int, float)) else str(value)


def _xlsx_records(content: bytes) -> list[Record]:
    sheet = load_workbook(BytesIO(content))["INVENTÁRIO"]
    return [
        (row[0], _bay(row[1]), _layer(row[2]), row[3], row[4])
        for row in sheet.iter_rows(min_row=5, values_only=True)
        if row[0] in {"DE", "EF"}
    ]


def _xls_records(content: bytes) -> list[Record]:
    sheet = open_workbook(file_contents=content).sheet_by_name("INVENTÁRIO")
    return [
        (sheet.cell_value(row, 0), _bay(sheet.cell_value(row, 1)), _layer(sheet.cell_value(row, 2)), sheet.cell_value(row, 3), int(sheet.cell_value(row, 4)))
        for row in range(4, sheet.nrows)
        if sheet.cell_value(row, 0) in {"DE", "EF"}
    ]


def _docx_table(document, first_header: str, third_header: str):
    return next(table for table in document.tables if table.rows[0].cells[0].text == first_header and len(table.rows[0].cells) > 2 and table.rows[0].cells[2].text == third_header)


def _docx_records(content: bytes) -> list[Record]:
    table = _docx_table(Document(BytesIO(content)), "Lado", "Camada")
    return [
        (cells[0], cells[1], _layer(cells[2]), cells[3], _int(cells[4]))
        for cells in ([cell.text for cell in row.cells] for row in table.rows[1:])
        if cells[0] in {"DE", "EF"}
    ]


def _pdf_lines(content: bytes) -> list[str]:
    return [line.strip() for page in PdfReader(BytesIO(content)).pages for line in (page.extract_text() or "").splitlines()]


def _pdf_section(lines: list[str], start: str, end: str | None) -> list[str]:
    first = lines.index(start)
    last = lines.index(end, first) if end else len(lines)
    return lines[first:last]


def _pdf_records(content: bytes) -> list[Record]:
    lines = _pdf_section(_pdf_lines(content), "3. Lançamentos por lado e vão", "4. Anexo: lotes consolidados")
    return [
        (lines[index], lines[index + 1], _layer(lines[index + 2]), lines[index + 3], _int(lines[index + 4]))
        for index in range(len(lines) - 4)
        if lines[index] in {"DE", "EF"} and LOT_PATTERN.match(lines[index + 3])
    ]


def _report_records(report: dict[str, object]) -> list[Record]:
    return [(row["side"], row["bay"], row["layer"], row["lot"], row["quantity"]) for row in report["records"]]  # type: ignore[index]


def test_fixture_is_the_real_inventory(report: dict[str, object]) -> None:
    assert len(FIXTURE["records"]) == EXPECTED["records"]
    assert sum(row["quantity"] for row in FIXTURE["records"]) == EXPECTED["pieces"]
    assert report["totalRecords"] == EXPECTED["records"]
    assert report["totalPieces"] == EXPECTED["pieces"]
    assert report["summary"]["lotsAnalyzed"] == EXPECTED["lots"]  # type: ignore[index]
    assert report["summary"]["lotsOk"] == EXPECTED["lotsOk"]  # type: ignore[index]
    assert report["summary"]["lotsForConference"] == EXPECTED["lotsForConference"]  # type: ignore[index]


def test_every_format_returns_exactly_the_input_records(report: dict[str, object], files: dict[str, bytes]) -> None:
    expected = _report_records(report)
    assert Counter(expected) == Counter(_input_records())
    extracted = {
        "xls": _xls_records(files["xls"]),
        "xlsx": _xlsx_records(files["xlsx"]),
        "pdf": _pdf_records(files["pdf"]),
        "docx": _docx_records(files["docx"]),
    }
    for name, records in extracted.items():
        assert records == expected, name


def test_totals_close_in_every_format(report: dict[str, object], files: dict[str, bytes]) -> None:
    view = build_report_view(report)
    assert sum(group["pieces"] for group in view.bay_groups) == EXPECTED["pieces"]
    assert len(view.bay_groups) == EXPECTED["bays"]
    assert sum(int(lot["totalQuantity"]) for lot in report["lots"]) == EXPECTED["pieces"]  # type: ignore[index]

    expected_summary = {
        "Total de registros": EXPECTED["records"],
        "Total de peças": EXPECTED["pieces"],
        "Total de lotes": EXPECTED["lots"],
        "Total de vãos": EXPECTED["bays"],
        "Lotes OK": EXPECTED["lotsOk"],
        "Lotes para conferência": EXPECTED["lotsForConference"],
    }
    xlsx_sheet = load_workbook(BytesIO(files["xlsx"]))["RESUMO"]
    xlsx_summary = {xlsx_sheet.cell(row, 1).value: xlsx_sheet.cell(row, 2).value for row in range(5, xlsx_sheet.max_row + 1)}
    xls_sheet = open_workbook(file_contents=files["xls"]).sheet_by_name("RESUMO")
    xls_summary = {xls_sheet.cell_value(row, 0): xls_sheet.cell_value(row, 1) for row in range(4, xls_sheet.nrows)}
    for label, value in expected_summary.items():
        assert xlsx_summary[label] == value, label
        assert xls_summary[label] == value, label

    # Painel do PDF e do Word: indicadores, OK × conferência e conferência dos totais.
    expected_kpis = {"Registros": "216", "Peças": "2.667", "Lotes": "207", "Vãos": "31"}
    document = Document(BytesIO(files["docx"]))
    kpi_table = next(table for table in document.tables if table.rows[0].cells[0].text.startswith("Registros\n"))
    assert dict(cell.text.split("\n") for cell in kpi_table.rows[0].cells) == expected_kpis
    panel = _pdf_section(_pdf_lines(files["pdf"]), "1. Painel", "2. Lotes para conferência")
    assert {label: panel[panel.index(label) + 1] for label in expected_kpis} == expected_kpis
    status = "200 lotes OK × 7 para conferência"
    assert status in panel
    assert status in "\n".join(paragraph.text for paragraph in document.paragraphs)

    expected_check = {
        "Soma dos lançamentos": "2.667",
        "Soma dos lotes": "2.667",
        "Total do resumo": "2.667",
        "Soma dos subtotais por vão": "2.667",
        "Registros: lançados × resumo": "216 × 216",
        "Resultado": "Confere",
    }
    assert {label: panel[panel.index(label) + 1] for label in expected_check} == expected_check
    check_table = _docx_nested_table(document, "Conferência dos totais")
    assert {row.cells[0].text: row.cells[1].text for row in check_table.rows[1:]} == expected_check

    inventory = load_workbook(BytesIO(files["xlsx"]))["INVENTÁRIO"]
    assert inventory.cell(inventory.max_row, 1).value == "TOTAL"
    assert inventory.cell(inventory.max_row, 5).value == EXPECTED["pieces"]


def _docx_nested_table(document, first_header: str):
    tables = list(document.tables) + [nested for table in document.tables for row in table.rows for cell in row.cells for nested in cell.tables]
    return next(table for table in tables if table.rows[0].cells[0].text == first_header)


def _lot_totals_xlsx(content: bytes, sheet_name: str) -> dict[str, int]:
    sheet = load_workbook(BytesIO(content))[sheet_name]
    return {row[0]: row[1] for row in sheet.iter_rows(min_row=5, values_only=True) if isinstance(row[0], str) and LOT_PATTERN.match(row[0])}


def _lot_totals_xls(content: bytes, sheet_name: str) -> dict[str, int]:
    sheet = open_workbook(file_contents=content).sheet_by_name(sheet_name)
    return {
        sheet.cell_value(row, 0): int(sheet.cell_value(row, 1))
        for row in range(4, sheet.nrows)
        if LOT_PATTERN.match(str(sheet.cell_value(row, 0)))
    }


def _lot_totals_docx(content: bytes, first_header: str, third_header: str, lot_column: int = 0) -> dict[str, int]:
    table = _docx_table(Document(BytesIO(content)), first_header, third_header)
    return {
        row.cells[lot_column].text: _int(row.cells[lot_column + 1].text)
        for row in table.rows[1:]
        if LOT_PATTERN.match(row.cells[lot_column].text)
    }


def _lot_totals_pdf(content: bytes, start: str, end: str | None) -> dict[str, int]:
    lines = _pdf_section(_pdf_lines(content), start, end)
    return {lines[index]: _int(lines[index + 1]) for index in range(len(lines) - 1) if LOT_PATTERN.match(lines[index])}


def test_lots_and_conference_match_across_formats(report: dict[str, object], files: dict[str, bytes]) -> None:
    lots = {str(lot["lot"]): int(lot["totalQuantity"]) for lot in report["lots"]}  # type: ignore[index]
    assert len(lots) == EXPECTED["lots"]
    assert _lot_totals_xlsx(files["xlsx"], "LOTES CONSOLIDADOS") == lots
    assert _lot_totals_xls(files["xls"], "LOTES CONSOLIDADOS") == lots
    assert _lot_totals_docx(files["docx"], "Lote", "Localização") == lots
    assert _lot_totals_pdf(files["pdf"], "4. Anexo: lotes consolidados", None) == lots

    conference = {
        str(lot["lot"]): int(lot["totalQuantity"])
        for lot in report["lots"]  # type: ignore[union-attr]
        if lot["presentation"]["requiresConference"]
    }
    assert len(conference) == EXPECTED["lotsForConference"]
    assert _lot_totals_xlsx(files["xlsx"], "LOTES PARA CONFERÊNCIA") == conference
    assert _lot_totals_xls(files["xls"], "LOTES PARA CONFERÊNCIA") == conference
    assert _lot_totals_docx(files["docx"], "Ordem", "Total", lot_column=1) == conference
    assert _lot_totals_pdf(files["pdf"], "2. Lotes para conferência", "3. Lançamentos por lado e vão") == conference


def test_excel_cell_types(files: dict[str, bytes]) -> None:
    modern = load_workbook(BytesIO(files["xlsx"]))
    legacy = open_workbook(file_contents=files["xls"], formatting_info=True)

    modern_summary = modern["RESUMO"]
    assert modern_summary.cell(5, 1).value == "Data do inventário"
    assert modern_summary.cell(5, 2).value == datetime(2026, 10, 6)
    assert modern_summary.cell(5, 2).number_format == "DD/MM/YYYY"
    assert modern_summary.cell(2, 1).value == "Revisão: 6 · Gerado em: 07/10/2026 às 20:20"
    legacy_summary = legacy.sheet_by_name("RESUMO")
    assert legacy_summary.cell_type(4, 1) == XL_CELL_DATE
    assert xldate_as_datetime(legacy_summary.cell_value(4, 1), legacy.datemode).date() == date(2026, 10, 6)
    labels = [legacy_summary.cell_value(row, 0) for row in range(4, legacy_summary.nrows)]
    assert labels.count("Data do inventário") == 1
    assert labels.count("Observações importantes") == 1

    modern_inventory = modern["INVENTÁRIO"]
    legacy_inventory = legacy.sheet_by_name("INVENTÁRIO")
    for row in range(5, modern_inventory.max_row):
        assert isinstance(modern_inventory.cell(row, 2).value, int)
        assert isinstance(modern_inventory.cell(row, 4).value, str) and modern_inventory.cell(row, 4).number_format == "@"
        assert isinstance(modern_inventory.cell(row, 5).value, int)
    for row in range(4, legacy_inventory.nrows - 1):
        assert legacy_inventory.cell_type(row, 1) == XL_CELL_NUMBER
        assert legacy_inventory.cell_type(row, 3) == XL_CELL_TEXT and len(legacy_inventory.cell_value(row, 3)) == 10
        assert legacy_inventory.cell_type(row, 4) == XL_CELL_NUMBER

    modern_conference = modern["LOTES PARA CONFERÊNCIA"]
    out_values = [modern_conference.cell(row, 6).value for row in range(5, modern_conference.max_row + 1)]
    assert all(value is None or isinstance(value, int) for value in out_values)
    assert None in out_values and 2 in out_values
    legacy_conference = legacy.sheet_by_name("LOTES PARA CONFERÊNCIA")
    assert {legacy_conference.cell_type(row, 5) for row in range(4, legacy_conference.nrows)} == {XL_CELL_NUMBER, XL_CELL_BLANK}

    assert modern["INVENTÁRIO"].oddFooter.center.text == "Página &P de &N"


def test_documents_show_local_time_dates_and_page_count(files: dict[str, bytes]) -> None:
    reader = PdfReader(BytesIO(files["pdf"]))
    first_page = reader.pages[0].extract_text()
    assert "Data do inventário: 06/10/2026 · Revisão: 6 · Gerado em: 07/10/2026 às 20:20" in first_page
    assert "+00:00" not in first_page and "2026-10-06" not in first_page
    total = len(reader.pages)
    for number, page in enumerate(reader.pages, start=1):
        assert f"Página {number} de {total}" in page.extract_text()

    document = Document(BytesIO(files["docx"]))
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Gerado em: 07/10/2026 às 20:20" in paragraphs
    footer_xml = document.sections[0].footer._element.xml
    assert "NUMPAGES" in footer_xml and " PAGE " in footer_xml
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "Data do inventário" not in cells
    assert "Observações importantes" not in cells


def test_pdf_has_no_orphan_heading_and_no_duplicated_observation(files: dict[str, bytes]) -> None:
    headings = set(SECTION_TITLES) | {"Quadro por lado", "Conferência dos totais", "Mapa de peças por lado e vão"}
    for page in PdfReader(BytesIO(files["pdf"])).pages:
        lines = [line.strip() for line in (page.extract_text() or "").splitlines() if line.strip()]
        content = [line for line in lines if not line.startswith("Página ")]
        assert content[-1] not in headings
    text = "\n".join(_pdf_lines(files["pdf"]))
    assert text.count("Há 7 lote(s) para conferência") == 1


def test_generated_at_uses_sao_paulo_time() -> None:
    assert format_generated_at(datetime(2026, 10, 7, 23, 30, tzinfo=timezone.utc)) == "07/10/2026 às 20:30"
    assert format_generated_at("2026-10-08T01:30:59.123456+00:00") == "07/10/2026 às 22:30"
    assert format_generated_at("2026-10-07T20:20:00-03:00") == "07/10/2026 às 20:20"


@pytest.mark.parametrize(
    "tamper",
    [
        lambda report: report.__setitem__("totalPieces", report["totalPieces"] + 1),
        lambda report: report.__setitem__("totalRecords", report["totalRecords"] - 1),
        lambda report: report["records"][0].__setitem__("quantity", report["records"][0]["quantity"] + 1),
        lambda report: report["lots"][0].__setitem__("totalQuantity", report["lots"][0]["totalQuantity"] + 1),
        lambda report: report["records"].pop(),
    ],
)
def test_inconsistent_report_never_generates_a_file(report: dict[str, object], tamper) -> None:
    broken = deepcopy(report)
    tamper(broken)
    for exporter in (export_xls, export_xlsx, export_pdf, export_docx):
        with pytest.raises(ReportConsistencyError, match="Relatório não gerado: os totais não conferem"):
            exporter(broken)


def test_export_endpoint_reports_the_inconsistency_clearly(report: dict[str, object]) -> None:
    from fastapi import HTTPException

    from app.main import _export_response

    broken = deepcopy(report)
    broken["totalPieces"] = int(broken["totalPieces"]) + 1  # type: ignore[call-overload]
    with pytest.raises(HTTPException) as error:
        _export_response(broken, "xls")
    assert error.value.status_code == 500
    assert "Relatório não gerado: os totais não conferem" in str(error.value.detail)
    assert "2.667" in str(error.value.detail) and "2.668" in str(error.value.detail)


def test_pdf_and_word_follow_the_delivery_structure(files: dict[str, bytes]) -> None:
    lines = _pdf_lines(files["pdf"])
    positions = [lines.index(title) for title in SECTION_TITLES[:4]]
    assert positions == sorted(positions)
    assert len(PdfReader(BytesIO(files["pdf"])).pages) <= 11

    document = Document(BytesIO(files["docx"]))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    word_positions = [paragraphs.index(title) for title in SECTION_TITLES[:4]]
    assert word_positions == sorted(word_positions)
    columns = [section._sectPr.find(qn("w:cols")) for section in document.sections]
    assert [element.get(qn("w:num")) if element is not None else None for element in columns][-1] == "2"


def test_conference_lots_follow_the_yard_route(report: dict[str, object], files: dict[str, bytes]) -> None:
    table = _docx_table(Document(BytesIO(files["docx"])), "Ordem", "Total")
    order = [row.cells[1].text for row in table.rows[1:]]
    assert order == ["2715284734", "2816141031", "2815286936", "2816112536", "2815458132", "2816112535", "2815602832"]
    assert [row.cells[0].text for row in table.rows[1:]] == [str(index) for index in range(1, 8)]


def test_grouped_records_have_exact_subtotals(report: dict[str, object], files: dict[str, bytes]) -> None:
    table = _docx_table(Document(BytesIO(files["docx"])), "Lado", "Camada")
    running = 0
    subtotals: list[int] = []
    side_totals: list[int] = []
    for row in table.rows[1:]:
        cells = [cell.text for cell in row.cells]
        if cells[0] in {"DE", "EF"}:
            running += _int(cells[4])
        elif cells[3].startswith("Subtotal"):
            assert _int(cells[4]) == running, cells[3]
            subtotals.append(running)
            running = 0
        elif cells[3].startswith("Total lado"):
            side_totals.append(_int(cells[4]))
        elif cells[0] == "TOTAL":
            assert _int(cells[4]) == EXPECTED["pieces"]
    assert len(subtotals) == EXPECTED["bays"]
    assert sum(subtotals) == sum(side_totals) == EXPECTED["pieces"]
    assert side_totals == [1484, 1183]


def test_color_only_marks_state(files: dict[str, bytes]) -> None:
    table = _docx_table(Document(BytesIO(files["docx"])), "Lote", "Localização")
    for row in table.rows[1:]:
        shading = row.cells[0]._tc.tcPr.find(qn("w:shd"))
        if row.cells[3].text == "OK":
            assert shading is None
        else:
            assert shading is not None
    header_fill = table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"))
    assert header_fill == "003B63"
