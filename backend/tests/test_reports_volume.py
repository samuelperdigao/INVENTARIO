from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from xlrd import open_workbook

from app.engine import AnalysisEntry
from app.reports import build_inventory_report_data, export_docx, export_pdf, export_xls, export_xlsx


def volume_entries() -> list[AnalysisEntry]:
    entries = [
        AnalysisEntry(side="EF", bay="15", layer="A1", lot="900001", quantity=19),
        AnalysisEntry(side="DE", bay="21", layer=None, lot="900001", quantity=1),
        AnalysisEntry(side="EF", bay="10", layer="A2", lot="900002", quantity=10),
        AnalysisEntry(side="DE", bay="10", layer=None, lot="900002", quantity=10),
        AnalysisEntry(side="EF", bay="11", layer=None, lot="900003", quantity=11),
        AnalysisEntry(side="DE", bay="11", layer="A3", lot="900003", quantity=9),
        AnalysisEntry(side="DE", bay="08", layer="A4", lot="900004", quantity=15),
        AnalysisEntry(side="EF", bay="12", layer=None, lot="900004", quantity=2),
        AnalysisEntry(side="DE", bay="20", layer="A5", lot="900004", quantity=3),
    ]
    for index in range(91):
        entries.append(
            AnalysisEntry(
                side="EF" if index % 2 == 0 else "DE",
                bay=str((index % 30) + 1),
                layer=None if index % 3 == 0 else f"A{(index % 10) + 1}",
                lot=f"{910000 + index:06d}",
                quantity=(index % 27) + 1,
            )
        )
    return entries


def volume_report() -> dict[str, object]:
    return build_inventory_report_data(
        "00000000-0000-4000-8000-000000000100",
        "2026-09-12",
        101,
        volume_entries(),
        datetime(2026, 9, 12, tzinfo=timezone.utc),
    )


def _xlsx_summary(workbook) -> dict[str, object]:
    sheet = workbook["RESUMO"]
    return {sheet.cell(row, 1).value: sheet.cell(row, 2).value for row in range(5, sheet.max_row + 1)}


def _xls_summary(workbook) -> dict[str, object]:
    sheet = workbook.sheet_by_name("RESUMO")
    return {sheet.cell_value(row, 0): sheet.cell_value(row, 1) for row in range(4, sheet.nrows)}


def test_100_records_generate_valid_equal_xls_and_xlsx_reports() -> None:
    report = volume_report()
    xls = export_xls(report)
    xlsx = export_xlsx(report)
    pdf = export_pdf(report)
    docx = export_docx(report)

    assert all(len(content) > 0 for content in (xls, xlsx, pdf, docx))
    assert xls[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    assert xlsx[:4] == b"PK\x03\x04"
    with ZipFile(BytesIO(xlsx)) as archive:
        assert archive.testzip() is None

    legacy = open_workbook(file_contents=xls)
    modern = load_workbook(BytesIO(xlsx))
    expected_tabs = ["RESUMO", "INVENTÁRIO", "LOTES CONSOLIDADOS", "DIVERGÊNCIAS"]
    assert legacy.sheet_names() == expected_tabs
    assert modern.sheetnames == expected_tabs
    assert sum(name.name == "_FilterDatabase" for name in legacy.name_obj_list) == 0
    assert all(not legacy.sheet_by_name(name).show_grid_lines for name in expected_tabs)
    assert all(modern[name].sheet_view.showGridLines is False for name in expected_tabs)
    assert all(modern[name].auto_filter.ref is None for name in expected_tabs)

    legacy_summary = _xls_summary(legacy)
    modern_summary = _xlsx_summary(modern)
    for label in ("Total de registros", "Total de peças", "Total de lotes", "Quantidade de lotes fragmentados"):
        assert legacy_summary[label] == modern_summary[label] == {
            "Total de registros": 100,
            "Total de peças": 1269,
            "Total de lotes": 95,
            "Quantidade de lotes fragmentados": 4,
        }[label]
    assert report["summary"]["fragmentedLots"] == 4

    legacy_inventory = legacy.sheet_by_name("INVENTÁRIO")
    modern_inventory = modern["INVENTÁRIO"]
    assert legacy_inventory.nrows == modern_inventory.max_row == 105
    assert legacy_inventory.cell_value(104, 4) == modern_inventory.cell(105, 5).value == 1269

    legacy_lots = legacy.sheet_by_name("LOTES CONSOLIDADOS")
    modern_lots = modern["LOTES CONSOLIDADOS"]
    legacy_lot_row = next(row for row in range(4, legacy_lots.nrows) if legacy_lots.cell_value(row, 0) == "900001")
    modern_lot_row = next(row for row in range(5, modern_lots.max_row + 1) if modern_lots.cell(row, 1).value == "900001")
    assert "PEÇA_SOLTEIRA" in str(legacy_lots.cell_value(legacy_lot_row, 5))
    assert "PEÇA_SOLTEIRA" in str(modern_lots.cell(modern_lot_row, 6).value)
    assert any(legacy_lots.cell_value(row, 0) == "910090" for row in range(4, legacy_lots.nrows))
    assert any(modern_lots.cell(row, 1).value == "910090" for row in range(5, modern_lots.max_row + 1))

    legacy_divergences = legacy.sheet_by_name("DIVERGÊNCIAS")
    modern_divergences = modern["DIVERGÊNCIAS"]
    assert any(legacy_divergences.cell_value(row, 0) == "900001" for row in range(4, legacy_divergences.nrows))
    assert any(modern_divergences.cell(row, 1).value == "900001" for row in range(5, modern_divergences.max_row + 1))

    pdf_reader = PdfReader(BytesIO(pdf))
    pdf_text = "".join(page.extract_text() or "" for page in pdf_reader.pages)
    assert pdf_reader.pages[0].mediabox.width > pdf_reader.pages[0].mediabox.height
    assert "900001" in pdf_text
    assert "910090" in pdf_text
    assert "PEÇA_SOLTEIRA" in pdf_text
    document = Document(BytesIO(docx))
    docx_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "900001" in docx_text
    assert "910090" in docx_text
    assert "PEÇA_SOLTEIRA" in docx_text
